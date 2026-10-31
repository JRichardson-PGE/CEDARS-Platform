import csv
import io
import zipfile
from datetime import date, timedelta
from decimal import Decimal

import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from django.utils import timezone

from cedars.core import storage
from cedars.modules.calculations import service as calculations
from cedars.modules.calculations.models import Run
from cedars.modules.claims import service, spec, upload
from cedars.modules.claims.models import ChangeLetter, ClaimProgramLine, ClaimYear
from cedars.services.accounts import service as access
from cedars.services.accounts.service import AccessDenied
from cedars.services.audit.models import AuditEvent
from cedars.services.registry import service as registry
from cedars.services.submissions import service as submissions
from cedars.services.submissions.models import SubmissionVersion

from .conftest import sign_in
from .test_budgets import release  # noqa: F401 - the engine release fixture
from .test_uploads import split

pytestmark = pytest.mark.django_db
DEFAULTS = {"text": "x", "text6": "x", "text255": "x", "number": "0", "integer": "2028", "bit": "0", "date": "2028-02-01"}
GIVEN = {
    "Claim": {"ClaimType": "Custom", "TotalFirstYearNetkWh": "800", "TotalFirstYearNetTherm": "0", "TotalGrossIncentive": "50"},
    "Measure": {"ClaimYearQuarter": "2028Q1", "NumUnits": "10", "UnitkWh1stBaseline": "100", "UnitEndUserRebate": "5", "NTGRkWh": "0.8",
                "MeasImpactType": "Cust", "IOU_AC_Territory": "PGE", "EUL_Yrs": "10", "RUL_Yrs": "0"},
    "ProgramCost": {"PrgYear": "2028", "ClaimYearQuarter": "2028Q1", "PA": "PGE", "AdminCostsOverheadAndGA": "1000", "EMV": "100"},
    "Site": {"SiteState": "CA", "SiteZipCode": "94111", "Residential_Flag": "1"},
}  # fmt: skip


def row(table: str, **values) -> dict:
    """A record with every required field filled in (type-appropriate defaults), then the values given."""
    out = {f.name: DEFAULTS[f.type] for f in spec.fields()[table] if f.required}
    out.update(GIVEN.get(table, {}))
    out.update(values)
    return out


def claim(cid: str, prg="PGE_Res", quarter="2028Q1", site="PGE-S1", **measure) -> dict:
    """A claim's Claim and Measure records (one each, as CEDARS has them)."""
    return {
        "Claim": [row("Claim", ClaimID=cid, SiteID=site)],
        "Measure": [row("Measure", ClaimID=cid, PrgID=prg, ClaimYearQuarter=quarter, **measure)],
    }


def tables(*parts, **more) -> dict:
    out: dict[str, list] = {}
    for p in parts + (more,):
        for t, rows in p.items():
            out.setdefault(t, []).extend(rows)
    return out


def claims_zip(t: dict, sep="|", names=None, full=False) -> bytes:
    """The tables as a zip of CSVs: the columns the records use (every specification column when `full`), in template order."""
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for name, rows in t.items():
            cols = [c for c in spec.columns(name) if full or any(c in r for r in rows)] + sorted(
                {c for r in rows for c in r} - set(spec.columns(name))
            )
            text = (
                sep.join((names or {}).get(c, c) for c in cols) + "\n" + "".join(sep.join(r.get(c, "") for c in cols) + "\n" for r in rows)
            )
            z.writestr(f"{name}.csv", text)
    return out.getvalue()


def one_claim(**costs) -> bytes:
    """PGE-2028-001 alone, with its program's cost row and site."""
    return claims_zip(
        tables(claim("PGE-2028-001"), ProgramCost=[row("ProgramCost", PrgID="PGE_Res", **costs)], Site=[row("Site", SiteID="PGE-S1")])
    )


def finish_run(run: Run, scale: float = 1.0) -> Run:
    """The worker's end, as a ReCET worker leaves it: group sums under out/_results/ (PGE_Res: TSB 4,500 x scale, 170 net kWh;
    PGE_Com: TSB 400) and a receipt."""
    rows = [
        {"_pa": "PGE", "_prg": "PGE_Res", "_prgname": "Res", "_sector": "Res", "_segment": "RA", "_year": "2028", "_n": 1,
         "TotalSystemBenefit": 3000.0 * scale, "TRCBen": 3000.0 * scale, "TRCCost": 2000.0, "PACBen": 3000.0 * scale, "PACCost": 1500.0,
         "FirstYearNetKWh": 100.0},
        {"_pa": "PGE", "_prg": "PGE_Res", "_prgname": "Res", "_sector": "Com", "_segment": "RA", "_year": "2028", "_n": 1,
         "TotalSystemBenefit": 1500.0 * scale, "TRCBen": 1500.0 * scale, "TRCCost": 1000.0, "PACBen": 1500.0 * scale, "PACCost": 500.0,
         "FirstYearNetKWh": 70.0},
        {"_pa": "PGE", "_prg": "PGE_Com", "_prgname": "Com", "_sector": "Com", "_segment": "RA", "_year": "2028", "_n": 1,
         "TotalSystemBenefit": 400.0, "TRCBen": 400.0, "TRCCost": 100.0, "PACBen": 400.0, "PACCost": 100.0, "FirstYearNetKWh": None},
    ]  # fmt: skip
    buf = io.BytesIO()
    pq.write_table(pa.Table.from_pylist(rows), buf)
    calculations.mark_started(run, worker="test")
    out = run.storage_prefix + "out/"
    storage.store("results").put_bytes(out + "_results/group_sums2_measure_level_run.parquet", buf.getvalue())
    storage.store("results").put_bytes(out + "receipt.json", b"{}")
    return calculations.finish(run, exit_code=0, receipt={"packs": {"files": {}}, "totals": {"TRCRatio": 1.5}})


def q1(full=False) -> bytes:
    return claims_zip(
        full=full,
        t=tables(
            claim("PGE-2028-001"),
            claim("PGE-2028-002", prg="PGE_Com", NumUnits="4"),
            ProgramCost=[row("ProgramCost", PrgID="PGE_Res"), row("ProgramCost", PrgID="PGE_Com", AdminCostsOverheadAndGA="300", EMV="0")],
            Site=[row("Site", SiteID="PGE-S1")],
            ContactClaim=[row("ContactClaim", ClaimID="PGE-2028-001", ContactID="PGE-C1")],
        ),
    )


@pytest.fixture
def programs(pge_admin, pge, cpuc_admin):
    sce = access.create_organisation(cpuc_admin, code="SCE", name="Southern California Edison", kind="pa", pa_type="IOU")
    registry.create_program(pge_admin, prg_id="PGE_Res", administrator=pge, valid_from=date(2024, 1, 1), name="Residential")
    registry.create_program(pge_admin, prg_id="PGE_Com", administrator=pge, valid_from=date(2024, 1, 1), name="Commercial")
    registry.create_program(cpuc_admin, prg_id="SCE_Ind", administrator=sce, valid_from=date(2024, 1, 1), name="Industrial")
    return sce


@pytest.fixture
def filer(make_user, pge, pge_admin):
    u = make_user("claims@pge.com", "Cly Mant")
    for p in ("claims.submit", "claims.confirm"):
        access.grant(pge_admin, u, pge, p)
    return u


@pytest.fixture
def year(cpuc_admin):
    now = timezone.now()
    due = {kind: now + timedelta(days=30 * (i + 1)) for i, (kind, _) in enumerate(service.PERIODS)}
    return service.add_year(cpuc_admin, year=2028, acc_version="2026", due=due)


def deadline(kind="Q1 claim"):
    return next(d for d in service.periods(2028) if d.kind == kind)


def confirmed(filer, pge, data, kind="Q1 claim", scale=1.0) -> SubmissionVersion:
    v = service.upload_claims(filer, deadline(kind), pge, data, "claims.zip")
    assert v.status == SubmissionVersion.Status.VALIDATED, service.findings(v)
    finish_run(Run.objects.get(pk=service.price(filer, v)), scale)
    v.refresh_from_db()
    return service.confirm(filer, v)


# --- the template and the upload


def test_the_template_is_cedars_claims_specification_with_the_platforms_own_tables():
    assert spec.TABLES == (
        "Claim", "Measure", "ProgramCost", "Site", "ContactClaim", "Participant", "InfrastructureCost", "ClaimSite", "Withdrawal",
    )  # fmt: skip
    assert len(spec.columns("Measure")) == 74 and "AC_Climate_Zone" in spec.columns("Measure")  # canonical names (decision L)
    assert spec.canonical("Measure", "E3ClimateZone") == "AC_Climate_Zone" and spec.canonical("Measure", "ceinputid") == "ClaimID"
    assert spec.field("InfrastructureCost", "TotalCost").type == "number" and spec.required("ContactClaim") == ["ClaimID", "ContactID"]
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(spec.template_xlsx()))
    assert wb.sheetnames == ["Fields", *spec.TABLES]
    assert [c.value for c in wb["Site"][1]] == spec.columns("Site")
    assert [c.value for c in wb["Withdrawal"][1]] == ["ClaimID", "WithdrawalReason"]  # WithdrawnIn is CEDARS' to fill in


def test_an_upload_is_a_zip_or_a_workbook_and_never_carries_pii():
    data = claims_zip(tables(claim("PGE-2028-001", Bogus="1")), names={"AC_Climate_Zone": "E3ClimateZone"})
    t = upload.read(data, "q1.zip")
    assert set(t) == {"Claim", "Measure"} and "AC_Climate_Zone" in t["Measure"].columns and t["Measure"].ignored == ["Bogus"]
    for extra, message in (("PII_Contact.csv", "personal information"), ("notes.docx", "only the claims tables")):
        z = io.BytesIO()
        with zipfile.ZipFile(z, "w") as f:
            f.writestr("Claim.csv", "ClaimID\nPGE-2028-1\n")
            f.writestr(extra, "x")
        with pytest.raises(upload.UploadError, match=message):
            upload.read(z.getvalue(), "q1.zip")
    with pytest.raises(upload.UploadError, match="not a valid .zip"):
        upload.read(b"ClaimID\n", "q1.zip")
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Fields"
    ws = wb.create_sheet("Site")
    ws.append(["SiteID", "SiteCity", "SiteState", "SiteZipCode"])
    ws.append(["PGE-S1", "Fresno", "CA", 93650])
    buf = io.BytesIO()
    wb.save(buf)
    t = upload.read(buf.getvalue(), "q1.xlsx")
    assert list(t) == ["Site"] and t["Site"].rows[0]["SiteZipCode"] == "93650"
    pii = Workbook()
    pii.active.title = "PII_Site"
    buf = io.BytesIO()
    pii.save(buf)
    with pytest.raises(upload.UploadError, match="personal information"):
        upload.read(buf.getvalue(), "q1.xlsx")


def test_only_a_cpuc_system_administrator_sets_up_claim_years(cpuc_admin, pge_admin, year):
    due = {k: timezone.now() for k, _ in service.PERIODS}
    with pytest.raises(AccessDenied, match="system administrator"):
        service.add_year(pge_admin, year=2029, acc_version="2026", due=due)
    with pytest.raises(AccessDenied, match="already set up"):
        service.add_year(cpuc_admin, year=2028, acc_version="2026", due=due)
    with pytest.raises(AccessDenied, match="due date"):
        service.add_year(cpuc_admin, year=2029, acc_version="2026", due={"Q1 claim": timezone.now()})
    assert [d.kind for d in service.periods(2028)] == [k for k, _ in service.PERIODS]
    assert [str(d.period) for d in service.periods(2028)] == ["2028Q1", "2028Q2", "2028Q3", "2028Q4", "2028"]
    assert AuditEvent.objects.filter(action="claim_year.set_up").exists()
    cy = service.change_setup(cpuc_admin, year, acc_version="2024", reason="ruling")
    assert cy.acc_version == "2024" and AuditEvent.objects.filter(action="claim_year.setup_changed").exists()


def test_the_default_acc_is_the_years_final_budget_filing():
    pytest.importorskip("recet.versions")
    assert service.default_acc(2028) == "2026" and service.default_acc(2029) == "2026"


# --- the first quarter


def test_a_first_quarter_gives_each_programs_year_to_date(filer, year, pge, programs):
    v = service.upload_claims(filer, deadline(), pge, q1(), "q1.zip")
    assert v.status == SubmissionVersion.Status.VALIDATED, service.findings(v)
    s = v.findings
    assert s["tables"] == {"Claim": 2, "Measure": 2, "ProgramCost": 2, "Site": 1, "ContactClaim": 1} and s["base"] is None
    assert s["compiled"] == s["tables"] and not s["replaced"]
    lines = {ln.prg_id: ln for ln in ClaimProgramLine.objects.filter(version=v)}
    res = lines["PGE_Res"]
    assert sum(Decimal(x) for x in res.program_costs.values()) == 1100 and res.claims == 1 and res.rebates_and_incentives == 50
    assert res.filed["TotalFirstYearNetkWh"] == 800 and lines["PGE_Com"].rebates_and_incentives == 20
    job = zipfile.ZipFile(io.BytesIO(storage.store("uploads").get_bytes(s["job_key"])))
    measure = list(csv.DictReader(io.StringIO(job.read("Measure.csv").decode())))
    assert [m["CEInputID"] for m in measure] == ["PGE-2028-001", "PGE-2028-002"] and "ClaimID" not in measure[0]
    t = service.table(v)
    assert t.total.expenditures == 1100 + 300 + 50 + 20 and not t.priced


def test_module_checks_reject_what_cedars_rejects(filer, year, pge, programs):
    bad = tables(
        claim("SCE-2028-001"),  # another PA's prefix
        claim("PGE-2028-002", quarter="2028Q2"),  # a later quarter than Q1
        claim("PGE-2028-003", prg="NOPE"),
        claim("PGE-2028-004", prg="SCE_Ind"),
        claim("PGE-2028-005", site="PGE-NOWHERE"),
        claim("PGE-2028-006", MeasImpactType="Cust-SEM"),  # no ParticipantID or SEM_Phase
        Measure=[row("Measure", ClaimID="PGE-2028-007", PrgID="PGE_Res")],  # no Claim record
        ProgramCost=[row("ProgramCost", PrgID="PGE_Res"), row("ProgramCost", PrgID="PGE_Res", PrgYear="2027")],
        Site=[row("Site", SiteID="PGE-S1"), row("Site", SiteID="PGE-S2", SiteState="NV", SiteZipCode="89501")],
        InfrastructureCost=[row("InfrastructureCost", ClaimID="PGE-2028-999", PrgID="PGE_Res")],
    )
    bad["Claim"].append(row("Claim", ClaimID="PGE-2028-008", SiteID="PGE-S1", NMEC_TU_Flag="1", ApplicationDate="soon"))
    v = service.upload_claims(filer, deadline(), pge, claims_zip(bad), "q1.zip")
    assert v.status == SubmissionVersion.Status.DRAFT
    found = {(f["table"], f["record_id"].split(" / ")[0], f["rule_id"]) for f in service.findings(v)}
    expected = {
        ("Claim", "SCE-2028-001", "CLM-ID"),
        ("Measure", "PGE-2028-002", "CLM-YEAR"),
        ("Measure", "PGE-2028-003", "CLM-PRG"),
        ("Measure", "PGE-2028-004", "CLM-ORG"),
        ("Claim", "PGE-2028-005", "CLM-REF"),
        ("Claim", "PGE-2028-006", "CLM-SEM"),
        ("Measure", "PGE-2028-007", "CLM-REF"),
        ("Claim", "PGE-2028-008", "CLM-REF"),  # its Measure is missing
        ("Claim", "PGE-2028-008", "CLM-NMEC"),
        ("Claim", "PGE-2028-008", "CLM-TYPE"),
        ("ProgramCost", "PGE_Res", "CLM-YEAR"),
        ("Site", "PGE-S2", "CLM-SITE"),
        ("InfrastructureCost", "PGE-2028-999", "CLM-REF"),
    }
    assert expected <= found, expected - found
    assert ("Measure", "PGE-2028-003", "CLM-REF") not in found  # an unknown program is reported once
    assert v.findings["rejected"] >= 12 and "CLM-ID" in service.findings_csv(v)
    with pytest.raises(AccessDenied, match="validated"):
        service.price(filer, v)


def test_a_new_upload_replaces_the_last_and_keeps_nothing_superseded(filer, year, pge, programs, django_capture_on_commit_callbacks):
    v = service.upload_claims(filer, deadline(), pge, q1(), "q1.zip")
    first = service.own_file(v)["key"]
    with django_capture_on_commit_callbacks(execute=True):
        v = service.upload_claims(filer, deadline(), pge, one_claim(), "q1b.zip")
    assert v.number == 1 and not storage.store("uploads").exists(first)
    assert set(ClaimProgramLine.objects.filter(version=v).values_list("prg_id", flat=True)) == {"PGE_Res"}


def test_a_claimer_needs_claims_submit(make_user, pge, pge_admin, year, programs):
    u = make_user("reader@pge.com")
    access.grant(pge_admin, u, pge, "claims.read")
    with pytest.raises(AccessDenied, match="claims.submit"):
        service.upload_claims(u, deadline(), pge, q1(), "q1.zip")


# --- priced, confirmed, published; the next quarter over it


def test_priced_then_confirmed_and_published(filer, year, pge, programs, release, make_user):  # noqa: F811
    v = service.upload_claims(filer, deadline(), pge, q1(), "q1.zip")
    run = Run.objects.get(pk=service.price(filer, v))
    assert run.acc_version == "2026" and run.settings == {"allocate_statewide": False, "acc_purpose": "claims"}
    with pytest.raises(AccessDenied, match="Price this version first"):
        service.confirm(filer, v)
    finish_run(run)
    v.refresh_from_db()
    assert service.priced(v)
    community = make_user("someone@example.com")
    assert not service.can_read(community, v)
    v = service.confirm(filer, v)
    assert v.status == SubmissionVersion.Status.PUBLISHED and service.can_read(community, v)
    assert [label for label, *_ in service.published_files(v)] == [
        "Year-to-date inputs (every claims table, compiled)",
        "This version's upload (its new and changed records)",
        "The job priced (Measure and ProgramCost)",
        "Run settings (job spec)",
        "Run receipt (engine, reference data, outputs)",
    ]
    t = service.table(v)
    assert t.priced and next(r for r in t.rows if r.prg_id == "PGE_Res").value("TotalSystemBenefit") == 4500


def test_a_later_quarter_is_compiled_over_the_confirmed_one(filer, year, pge, programs, release):  # noqa: F811
    first = confirmed(filer, pge, q1())
    q2 = tables(
        claim("PGE-2028-001", NumUnits="20"),  # a correction of a Q1 claim
        claim("PGE-2028-010", quarter="2028Q2"),
        ProgramCost=[row("ProgramCost", PrgID="PGE_Res", ClaimYearQuarter="2028Q2", AdminCostsOverheadAndGA="2500")],
        ContactClaim=[row("ContactClaim", ClaimID="PGE-2028-001", ContactID="PGE-C2")],
    )
    v = service.upload_claims(filer, deadline("Q2 claim"), pge, claims_zip(q2), "q2.zip")
    assert v.status == SubmissionVersion.Status.VALIDATED, service.findings(v)
    s = v.findings
    assert s["base"] == first.pk and s["changed_claims"] == ["pge-2028-001"] and s["changed_claims_count"] == 1
    assert s["replaced"] == {"Claim": 1, "Measure": 1, "ProgramCost": 1, "ContactClaim": 1}
    assert s["compiled"] == {"Claim": 3, "Measure": 3, "ProgramCost": 2, "Site": 1, "ContactClaim": 1}
    compiled = service.compiled_tables(v)
    assert {r["ContactID"] for r in compiled["ContactClaim"].rows} == {"PGE-C2"}  # a claim's contacts are replaced as a set
    lines = {ln.prg_id: ln for ln in v.claim_lines.all()}
    assert lines["PGE_Res"].claims == 2 and lines["PGE_Res"].rebates_and_incentives == 20 * 5 + 10 * 5
    assert sum(Decimal(x) for x in lines["PGE_Res"].program_costs.values()) == 2600  # the year-to-date ProgramCost row
    assert lines["PGE_Com"].claims == 1  # Q1's claim carries over


def test_quarters_are_confirmed_in_order(filer, year, pge, programs, release, cpuc_admin, signed_in):  # noqa: F811
    v1 = service.upload_claims(filer, deadline(), pge, q1(), "q1.zip")
    with pytest.raises(AccessDenied, match="Confirm the 2028 Q1 claim first"):
        service.upload_claims(filer, deadline("Q2 claim"), pge, q1(), "q2.zip")
    page = signed_in(filer).get(f"/claims/periods/{deadline('Q2 claim').pk}/").content.decode()  # the page says so before any form
    assert "Confirm the 2028 Q1 claim first" in page and 'name="file"' not in page
    finish_run(Run.objects.get(pk=service.price(filer, v1)))
    v1.refresh_from_db()
    v1 = service.confirm(filer, v1)
    v2 = service.upload_claims(filer, deadline("Q2 claim"), pge, claims_zip(claim("PGE-2028-010", quarter="2028Q2")), "q2.zip")
    assert v2.findings["base"] == v1.pk and not service.stale(v2)
    # Q1 un-confirmed and changed under Q2: Q2 must be checked again before it can be confirmed
    submissions.unconfirm(filer, v1, "fix a cost")
    v2.refresh_from_db()
    assert service.stale(v2)
    v1 = service.upload_claims(filer, deadline(), pge, one_claim(EMV="5"), "q1.zip")
    finish_run(Run.objects.get(pk=service.price(filer, v1)))
    v1.refresh_from_db()
    v1 = service.confirm(filer, v1)
    v2.refresh_from_db()
    assert service.stale(v2)  # same Q1 version, different records
    finish_run(Run.objects.get(pk=service.price(filer, v2)))
    v2.refresh_from_db()
    with pytest.raises(AccessDenied, match="Check it again"):
        service.confirm(filer, v2)
    v2 = service.check(filer, v2)
    assert not service.stale(v2) and not service.priced(v2)  # compiled again: the year to date changed, so it is priced again
    finish_run(Run.objects.get(pk=service.price(filer, v2)))
    v2.refresh_from_db()
    v2 = service.confirm(filer, v2)
    # Q2 stands on Q1: Q1 cannot take uploads, be un-confirmed, or be reopened after its due date
    with pytest.raises(AccessDenied, match="already confirmed"):
        service.upload_claims(filer, deadline(), pge, q1(), "q1.zip")
    with pytest.raises(AccessDenied, match="confirmed on top of this one"):
        service.unconfirm(filer, v1, "again")
    d = deadline()
    submissions.move_deadline(cpuc_admin, d, timezone.now() - timedelta(days=1), "test")
    with pytest.raises(AccessDenied, match="most recently confirmed"):
        service.request_reopen(filer, v1.submission, "late correction")


def test_the_claims_totals_hook_gives_the_year_to_date_by_program(filer, year, pge, programs, release):  # noqa: F811
    assert submissions.claims_totals(pge, 2028) == []
    confirmed(filer, pge, q1())
    rows = {r["prg_id"]: r for r in submissions.claims_totals(pge, 2028)}
    res = rows["PGE_Res"]
    assert res["quarters"] == [1] and res["final"] is False and res["records"] == 1 and res["pa"] == "PGE"
    assert res["expenditures"] == 1100 + 50 and res["incentives"] == 50 and res["tsb"] == 4500 and res["net_kwh"] == 170
    assert rows["PGE_Com"]["expenditures"] == 300 + 20


def test_a_reopened_version_is_compiled_from_its_predecessors_upload(
    filer,
    year,
    pge,
    programs,
    release,  # noqa: F811
    cpuc_admin,
    django_capture_on_commit_callbacks,
):
    v1 = confirmed(filer, pge, q1())
    submissions.move_deadline(cpuc_admin, deadline(), timezone.now() - timedelta(days=1), "test")
    submissions.set_policy(cpuc_admin, "claims", reopen_after_due="open")
    v2 = service.request_reopen(filer, v1.submission, "a missed claim")
    assert v2.number == 2 and not v2.findings
    service.ensure_checked(v2)
    v2.refresh_from_db()
    assert v2.status == SubmissionVersion.Status.VALIDATED and v2.findings["compiled"] == v1.findings["compiled"]
    # the copy is v2's own: v1's published files are never written over or deleted
    v1_files = [key for _, area, key in service.published_files(v1) if area == "uploads"]
    assert service.own_file(v2)["key"].startswith(f"claims/{v2.submission_id}/v2/") and v2.findings["compiled_key"] not in v1_files
    with django_capture_on_commit_callbacks(execute=True):
        service.upload_claims(filer, deadline(), pge, one_claim(EMV="7"), "correction.zip")
    assert all(storage.store("uploads").exists(key) for key in v1_files)


# --- withdrawals and corrections


def withdraw(*claims: str, reason="Customer canceled the project") -> dict:
    return {"Withdrawal": [row("Withdrawal", ClaimID=c, WithdrawalReason=reason) for c in claims]}


def test_a_withdrawn_claim_keeps_its_records_but_is_not_priced_or_counted(filer, year, pge, programs, release):  # noqa: F811
    confirmed(filer, pge, q1())
    q2 = tables(claim("PGE-2028-010", quarter="2028Q2"), withdraw("pge-2028-002"))
    q2["Withdrawal"][0]["WithdrawnIn"] = "made up"  # CEDARS fills it in, not the PA
    v = service.upload_claims(filer, deadline("Q2 claim"), pge, claims_zip(q2), "q2.zip")
    assert v.status == SubmissionVersion.Status.VALIDATED, service.findings(v)
    assert v.findings["withdrawn"] == 1 and v.findings["withdrawn_total"] == 1
    assert service.own_file(v)["ignored_columns"] == {"Withdrawal": ["WithdrawnIn"]}
    compiled = service.compiled_tables(v)
    assert "PGE-2028-002" in {r["ClaimID"] for r in compiled["Measure"].rows}  # kept in the year's history
    assert compiled["Withdrawal"].rows == [
        {"ClaimID": "pge-2028-002", "WithdrawalReason": "Customer canceled the project", "WithdrawnIn": "2028 Q2 claim, version 1"}
    ]
    job = zipfile.ZipFile(io.BytesIO(storage.store("uploads").get_bytes(v.findings["job_key"])))
    assert [m["CEInputID"] for m in csv.DictReader(io.StringIO(job.read("Measure.csv").decode()))] == ["PGE-2028-001", "PGE-2028-010"]
    lines = {ln.prg_id: ln for ln in v.claim_lines.all()}
    assert lines["PGE_Com"].claims == 0 and lines["PGE_Com"].rebates_and_incentives == 0 and lines["PGE_Com"].filed == {}
    assert [(w["claim"], w["confirmed_at"]) for w in service.withdrawals(v)] == [("pge-2028-002", None)]
    finish_run(Run.objects.get(pk=service.price(filer, v)))
    v.refresh_from_db()
    v = service.confirm(filer, v)
    assert service.withdrawals(v)[0]["confirmed_at"] == v.confirmed_at

    # a claim is withdrawn once, only after it is confirmed, with its records left as filed; its ClaimID is never filed again
    bad = tables(
        claim("PGE-2028-002", quarter="2028Q3"),
        claim("PGE-2028-020", quarter="2028Q3"),
        withdraw("PGE-2028-002", "PGE-2028-020", "PGE-2028-404"),
    )
    v3 = service.upload_claims(filer, deadline("Q3 claim"), pge, claims_zip(bad), "q3.zip")
    found = {(f["table"], f["record_id"], f["message"].split(" ")[-1]) for f in service.findings(v3) if f["rule_id"] == "CLM-WD"}
    assert found == {
        ("Withdrawal", "PGE-2028-002", "1"),  # already withdrawn, in the Q2 claim, version 1
        ("Withdrawal", "PGE-2028-020", "withdrawn"),  # not confirmed before: leave it out instead
        ("Withdrawal", "PGE-2028-404", "withdrawn"),
        ("Claim", "PGE-2028-020", "it"),  # withdrawn and filed in one upload
        ("Measure", "PGE-2028-020", "it"),
        ("Claim", "PGE-2028-002", "it"),  # withdrawn and filed in one upload, and withdrawn before
        ("Measure", "PGE-2028-002", "it"),
    }, found
    v3 = service.upload_claims(filer, deadline("Q3 claim"), pge, claims_zip(claim("PGE-2028-002", quarter="2028Q3")), "q3.zip")
    assert {f["message"] for f in service.findings(v3) if f["rule_id"] == "CLM-WD"} == {
        "Claim PGE-2028-002 was withdrawn in the 2028 Q2 claim, version 1"
    }


def test_a_correction_publishes_a_change_letter(
    client,
    filer,
    year,
    pge,
    programs,
    release,  # noqa: F811
    cpuc_admin,
    make_user,
    django_capture_on_commit_callbacks,
):
    v1 = confirmed(filer, pge, q1())
    submissions.move_deadline(cpuc_admin, deadline(), timezone.now() - timedelta(days=1), "test")
    submissions.set_policy(cpuc_admin, "claims", reopen_after_due="open")
    v2 = service.request_reopen(filer, v1.submission, "PGE-2028-002 was canceled; PGE-2028-001 had 20 units")
    correction = tables(
        claim("PGE-2028-001", NumUnits="20"),
        claim("PGE-2028-003"),
        withdraw("PGE-2028-002"),
        ProgramCost=[row("ProgramCost", PrgID="PGE_Res"), row("ProgramCost", PrgID="PGE_Com", AdminCostsOverheadAndGA="300", EMV="0")],
        Site=[row("Site", SiteID="PGE-S1")],
        ContactClaim=[row("ContactClaim", ClaimID="PGE-2028-001", ContactID="PGE-C1")],
    )
    with django_capture_on_commit_callbacks(execute=True):
        v2 = service.upload_claims(filer, deadline(), pge, claims_zip(correction), "correction.zip")
    assert v2.status == SubmissionVersion.Status.VALIDATED, service.findings(v2)
    compiled = service.compiled_tables(v2)
    assert {r["ClaimID"] for r in compiled["Measure"].rows} == {"PGE-2028-001", "PGE-2028-002", "PGE-2028-003"}  # 002 as filed
    assert not ChangeLetter.objects.exists()
    finish_run(Run.objects.get(pk=service.price(filer, v2)), scale=2.0)
    v2.refresh_from_db()
    v2 = service.confirm(filer, v2)
    letter = ChangeLetter.objects.get()
    assert letter.version == v2 and letter.previous == v1 and letter.reason.startswith("PGE-2028-002 was canceled")
    c = letter.changes
    assert c["previous_label"] == "2028 Q1 claim, version 1"
    assert {k: g["count"] for k, g in c["claims"].items()} == {"added": 1, "changed": 1, "removed": 0, "reinstated": 0, "withdrawn": 1}
    assert c["claims"]["added"]["ids"] == ["PGE-2028-003"] and c["claims"]["changed"]["ids"] == ["PGE-2028-001"]
    assert c["claims"]["withdrawn"]["ids"] == [{"claim": "PGE-2028-002", "reason": "Customer canceled the project"}]
    by_table = {t["table"]: t for t in c["tables"]}
    assert by_table["Measure"] == {"table": "Measure", "before": 2, "after": 3, "added": 1, "changed": 1, "removed": 0}
    assert by_table["Withdrawal"]["added"] == 1
    programs_ = {p["prg_id"]: p for p in c["programs"]}
    assert programs_["PGE_Res"]["before"]["claims"] == 1 and programs_["PGE_Res"]["after"]["claims"] == 2
    assert programs_["PGE_Res"]["before"]["TotalSystemBenefit"] == 4500 and programs_["PGE_Res"]["after"]["TotalSystemBenefit"] == 9000
    assert programs_["PGE_Com"]["after"]["claims"] == 0 and programs_["PGE_Com"]["after"]["expenditures"] == 300
    assert c["totals"]["before"]["expenditures"] == 1470 and c["totals"]["after"]["expenditures"] == 1100 + 20 * 5 + 10 * 5 + 300
    event = AuditEvent.objects.get(action="claims.change_letter_published")
    assert event.after["withdrawn"] == 1 and event.reason == letter.reason
    # the letter is append-only, and the version it replaces stays as confirmed
    from django.db import DatabaseError, transaction

    with pytest.raises(DatabaseError, match="append-only"), transaction.atomic():
        ChangeLetter.objects.filter(pk=letter.pk).update(reason="edited")
    v1.refresh_from_db()
    assert v1.status == SubmissionVersion.Status.PUBLISHED and service.standing(pge, deadline()) == v2
    # the change log, the letter and the linked-claims lookup are the public record of corrections (BR-09.3): no sign-in needed
    for path in ("/claims/changes/", f"/claims/changes/{letter.pk}/", "/claims/links/?kind=claim&q=PGE-2028-001"):
        r = client.get(path)
        assert r.status_code == 200, path
    assert "PGE-2028-002 was canceled" in client.get(f"/claims/changes/{letter.pk}/").content.decode()

    # pages: the change log and the letter are open to every signed-in user, linked from the versions
    sign_in(client, make_user("someone@example.com"))
    r = client.get(reverse("claims:changes"))
    assert r.status_code == 200 and b"PGE 2028 Q1 claim, version 2" in r.content
    assert client.get(reverse("claims:changes") + "?org=PGE&year=2027").content.count(b"version 2</a>") == 0
    r = client.get(reverse("claims:change_letter", args=[letter.pk]))
    assert r.status_code == 200 and b"Customer canceled the project" in r.content and b"PGE-2028-003" in r.content
    assert b"+$4,500" in r.content  # PGE_Res's total system benefit
    r = client.get(reverse("claims:version", args=[v2.pk]))
    assert b"change letter" in r.content and b"Withdrawn claims" in r.content and b"2028 Q1 claim, version 2" in r.content


# --- links between claims


def linked(cid: str, quarter="2028Q1", impact="Cust-NMEC", parent="", **claim_fields) -> dict:
    """A claim with the Claim fields that link it: a NMEC original (no parent) or true-up (a parent), or any other claim."""
    t = claim(cid, quarter=quarter, MeasImpactType=impact)
    if "NMEC" in impact:
        claim_fields.setdefault("NMEC_TU_Flag", "1" if parent else "0")
    t["Claim"][0].update(ParentClaimID=parent, **claim_fields)
    return t


def test_nmec_chains_projects_and_sem_participants_link_up(client, filer, year, pge, programs, release, cpuc_admin, make_user):  # noqa: F811
    cost = {"ProgramCost": [row("ProgramCost", PrgID="PGE_Res")], "Site": [row("Site", SiteID="PGE-S1"), row("Site", SiteID="PGE-S2")]}
    first = tables(
        linked("PGE-2028-001", ProjectID="PGE-P1"),
        linked("PGE-2028-002", impact="Cust", ProjectID="PGE-P1"),
        linked("PGE-2028-003", impact="Cust-SEM", ParticipantID="PGE-SEM-7", SEM_Phase="C1-RP1-2028"),
        cost,
    )
    confirmed(filer, pge, claims_zip(first))

    # the chain's checks: a parent the PA filed, not withdrawn, not the claim itself or one of its true-ups; a non-NMEC parent warns
    bad = tables(
        linked("PGE-2028-012", "2028Q2", parent="PGE-2028-999"),
        linked("PGE-2028-013", "2028Q2", parent="PGE-2028-013"),
        linked("PGE-2028-014", "2028Q2", parent="PGE-2028-003"),
        linked("PGE-2028-015", "2028Q2", parent="PGE-2028-016"),
        linked("PGE-2028-016", "2028Q2", parent="PGE-2028-015"),
        linked("PGE-2028-017", "2028Q2", parent="PGE-2028-002"),
        withdraw("PGE-2028-003"),
        ClaimSite=[row("ClaimSite", ClaimID="PGE-2028-012", SiteID="PGE-S9")],
    )
    v = service.upload_claims(filer, deadline("Q2 claim"), pge, claims_zip(bad), "q2.zip")
    found = {(f["record_id"].split(" / ")[0], f["severity"], f["message"].split(" ")[-1]) for f in service.findings(v)
             if f["rule_id"] in ("CLM-CHAIN", "CLM-REF") and f["table"] in ("Claim", "ClaimSite")}  # fmt: skip
    assert found == {
        ("PGE-2028-012", "reject", "filed"),
        ("PGE-2028-013", "reject", "itself"),
        ("PGE-2028-014", "reject", "claim"),  # withdrawn by this upload
        ("PGE-2028-015", "reject", "true-ups"),
        ("PGE-2028-016", "reject", "true-ups"),
        ("PGE-2028-017", "warning", "(Cust)"),
        ("PGE-2028-012", "reject", "record"),  # a ClaimSite's site must have a Site record
    }, found

    # Q2: a true-up of the Q1 NMEC claim covering two sites, another meter of the project, the SEM participant's next phase
    second = tables(
        linked("PGE-2028-010", "2028Q2", parent="PGE-2028-001", ProjectID="PGE-P1", TotalFirstYearNetkWh="1000"),
        linked("PGE-2028-011", "2028Q2", impact="Cust-SEM", ParticipantID="pge-sem-7", SEM_Phase="C1-RP2-2028"),
        ClaimSite=[row("ClaimSite", ClaimID="PGE-2028-010", SiteID=s) for s in ("PGE-S1", "PGE-S2")],
    )
    v2 = confirmed(filer, pge, claims_zip(second), kind="Q2 claim")
    rec = v2.claim_records.get(claim_id="PGE-2028-010")
    assert rec.true_up and rec.parent_claim_id == "PGE-2028-001" and rec.sites == ["PGE-S1", "PGE-S2"] and rec.project_id == "PGE-P1"
    links = service.chain("pge-2028-010")
    assert [(link.record.claim_id, link.depth) for link in links] == [("PGE-2028-001", 0), ("PGE-2028-010", 1)]
    assert links[1].change["TotalFirstYearNetkWh"] == 200 and links[0].record.version == v2  # the year's standing version
    assert [r.claim_id for r in service.project_claims("pge-p1")] == ["PGE-2028-001", "PGE-2028-002", "PGE-2028-010"]
    assert [r.sem_phase for r in service.participant_claims("PGE-SEM-7")] == ["C1-RP1-2028", "C1-RP2-2028"]
    assert "PGE-2028-010,PGE_Res,PGE-P1,PGE-2028-001,1,Cust-NMEC,,,PGE-S1,PGE-S1;PGE-S2,active" in service.links_csv(v2)

    # a later year's true-up names a confirmed earlier year's claim; the chain runs across the years
    due = {k: timezone.now() + timedelta(days=400) for k, _ in service.PERIODS}
    service.add_year(cpuc_admin, year=2029, acc_version="2026", due=due)
    d29 = next(d for d in service.periods(2029) if d.kind == "Q1 claim")
    nxt = tables(
        linked("PGE-2029-001", "2029Q1", parent="PGE-2028-010", ProjectID="PGE-P1"),
        linked("PGE-2029-002", "2029Q1", parent="PGE-2028-404"),
        ProgramCost=[row("ProgramCost", PrgID="PGE_Res", PrgYear="2029", ClaimYearQuarter="2029Q1")],
        Site=[row("Site", SiteID="PGE-S1")],
    )
    v29 = service.upload_claims(filer, d29, pge, claims_zip(nxt), "2029q1.zip")
    assert {(f["record_id"], f["rule_id"]) for f in service.findings(v29) if f["severity"] == "reject"} == {("PGE-2029-002", "CLM-CHAIN")}
    nxt = {t: [r for r in rows if r.get("ClaimID") != "PGE-2029-002"] for t, rows in nxt.items()}
    v29 = service.upload_claims(filer, d29, pge, claims_zip(nxt), "2029q1.zip")
    finish_run(Run.objects.get(pk=service.price(filer, v29)))
    v29.refresh_from_db()
    service.confirm(filer, v29)
    assert [link.record.claim_id for link in service.chain("PGE-2028-001")] == ["PGE-2028-001", "PGE-2028-010", "PGE-2029-001"]

    # the page, open to every signed-in user
    sign_in(client, make_user("someone@example.com"))
    r = client.get(reverse("claims:links") + "?kind=claim&q=PGE-2029-001")
    assert r.status_code == 200 and b"True-up 2" in r.content and b"Original claim" in r.content and b"+200" in r.content
    r = client.get(reverse("claims:links") + "?kind=project&q=PGE-P1")
    assert r.content.count(b"/claims/versions/") == 4
    assert b"No confirmed claim matches" in client.get(reverse("claims:links") + "?kind=participant&q=nobody").content
    r = client.get(reverse("claims:links_csv", args=[v2.pk]))
    assert r.status_code == 200 and b"PGE-SEM-7" in r.content


# --- the claims rule set


def test_the_claims_rules_run_on_measure_and_program_cost(filer, year, pge, programs, cpuc_admin):
    pytest.importorskip("recet.filing_rules", reason="the rule library runs on the pinned ReCET engine")
    from cedars.services.rules import service as rules

    rs = rules.import_from_engine(cpuc_admin, "claims")
    rules.publish(cpuc_admin, rs)
    v = service.upload_claims(filer, deadline(), pge, q1(full=True), "q1.zip")
    found = service.findings(v)
    assert v.findings["rule_set"].startswith(rs.label)
    assert {f["rule_id"] for f in found if f["table"] in ("Measure", "ProgramCost")} - {
        "CLM-REF",
        "CLM-PRG",
        "CLM-ORG",
        "CLM-YEAR",
        "CLM-ID",
    }
    assert not [f for f in found if f["rule_id"] in ("CLM-REQ", "CLM-TYPE") and f["table"] in ("Measure", "ProgramCost")]
    assert {f["record_id"] for f in found if f["table"] == "Measure"} <= {"PGE-2028-001", "PGE-2028-002"}  # keyed by ClaimID
    # a Measure file without the specification's columns is rejected on them before its rules run
    v = service.upload_claims(filer, deadline(), pge, q1(), "q1.zip")
    assert {f["rule_id"] for f in service.findings(v) if f["table"] == "Measure" and f["severity"] == "reject"} == {"CLM-COL"}


# --- pages


def test_the_pages(client, filer, year, pge, programs, release, cpuc_admin, make_user):  # noqa: F811
    v = confirmed(filer, pge, q1())
    d2 = deadline("Q2 claim")
    draft = service.upload_claims(filer, d2, pge, claims_zip(claim("SCE-2028-1")), "bad.zip")
    sign_in(client, filer)
    r = client.get(reverse("claims:index"))
    assert r.status_code == 200 and b"2028 claims" in r.content
    assert b">Published</span>" in r.content and b"version 1" in r.content  # the status badge and the version beside it
    r = client.get(reverse("claims:period", args=[deadline().pk]))
    assert b"Year-to-date results" in r.content and b"Records by table" in r.content
    r = client.get(reverse("claims:period", args=[d2.pk]))
    assert b"CLM-ID" in r.content and b"Earlier confirmations this year" in r.content and b"year-to-date inputs" in r.content
    r = client.get(reverse("claims:version", args=[v.pk]))
    assert r.status_code == 200 and b"PGE_Res" in r.content and b"Filed first-year net kWh" in r.content
    r = client.get(reverse("claims:version_csv", args=[v.pk]))
    rows = list(csv.DictReader(io.StringIO(r.content.decode())))
    assert rows[-1]["PrgID"] == "Total" and float(rows[-1]["Expenditures"]) == 1470 and float(rows[-1]["TotalFirstYearNetkWh"]) == 1600
    assert b"CLM-ID" in client.get(reverse("claims:findings_csv", args=[draft.pk])).content
    r = client.get(reverse("claims:file", args=[v.pk, 0]))
    assert r.status_code == 200 and zipfile.ZipFile(io.BytesIO(b"".join(r.streaming_content))).namelist()[0] == "Claim.csv"
    assert client.get(reverse("claims:template")).status_code == 200
    community = make_user("someone@example.com")
    client.logout()
    sign_in(client, community)
    assert client.get(reverse("claims:version", args=[v.pk])).status_code == 200  # published: public to signed-in users
    assert client.get(reverse("claims:version", args=[draft.pk])).status_code == 404
    assert client.get(reverse("claims:year_new")).status_code == 404
    client.logout()
    sign_in(client, cpuc_admin)
    form = {"year": 2029, "acc_version": "2026", "notes": ""} | {f"due_{k.split()[0].lower()}": "2029-05-01" for k, _ in service.PERIODS}
    r = client.post(reverse("claims:year_new"), form)
    assert r.status_code == 302 and ClaimYear.objects.filter(year=2029).exists()
    r = client.post(reverse("claims:setup", args=[2029]), {"acc_version": "2024", "reason": "ruling"})
    assert r.status_code == 302 and ClaimYear.objects.get(year=2029).acc_version == "2024"
    r = client.get(reverse("claims:period", args=[deadline().pk]) + "?org=PGE")  # the CPUC reads any filer's confirmations
    assert r.status_code == 200 and b"Year-to-date results" in r.content


def test_uploading_through_the_page(client, filer, year, pge, programs):
    sign_in(client, filer)
    from django.core.files.uploadedfile import SimpleUploadedFile

    r = client.post(reverse("claims:upload", args=[deadline().pk]), {"file": SimpleUploadedFile("q1.zip", q1())}, follow=True)
    assert b"Checked 8 records: none rejected" in r.content
    r = client.post(reverse("claims:upload", args=[deadline().pk]), {"file": SimpleUploadedFile("q1.zip", b"nope")}, follow=True)
    assert b"not a valid .zip" in r.content


def test_claims_in_several_files_and_a_pii_table_among_them(client, filer, year, pge, programs):
    sign_in(client, filer)
    url = reverse("claims:upload", args=[deadline().pk])
    files = split(q1())
    r = client.post(url, {"file": files}, follow=True)
    assert b"Checked 8 records: none rejected" in r.content
    v = SubmissionVersion.objects.get()
    assert [s["name"] for s in v.files[0]["sources"]] == [f.name for f in files]
    pii = SimpleUploadedFile("PII_Contact_2.csv", b"ContactID\nPGE-C1\n")
    r = client.post(url, {"file": [*split(q1()), pii]}, follow=True)
    assert b"holds personal information" in r.content
