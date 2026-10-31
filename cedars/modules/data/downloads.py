"""Stamped downloads (BR-07.2): a query's rows as CSV, XLSX or Parquet, with the release, its as-of time, the snapshots read, the
filters and exclusions, and each metric's canonical fields, so a file can be traced back to exactly what it shows.

Columns: the dimensions, each metric as shown (scaled, its unit in the heading), then each summed canonical field under its own name
and the number of records. CSV carries the stamp as "#" lines above the header; XLSX on an About sheet; Parquet in the schema's
"cedars" metadata."""

import csv
import io
import json

from django.utils import timezone

from .query import Result

FORMATS = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "parquet": "application/vnd.apache.parquet",
}


def table(result: Result, totals: bool = True) -> tuple[list[str], list[list]]:
    """The download's header and rows, the totals row last when the rows are grouped (not in Parquet, which keeps one type per column)."""
    header = [d.label for d in result.dimension_list]
    header += [f"{m.label} ({m.unit})" if m.unit else m.label for m in result.metric_list]
    header += result.fields + ["Records"]
    rows = [_row(result, r) for r in result.rows]
    if totals and result.query.by and result.rows:
        total = dict(result.totals)
        rows.append(_row(result, total, label="Total"))
    return header, rows


def _row(result: Result, r: dict, label: str | None = None) -> list:
    dims = [label if label and i == 0 else (None if label else r.get(d.code)) for i, d in enumerate(result.dimension_list)]
    return dims + [r.get(m.code) for m in result.metric_list] + [r["_sums"].get(f) for f in result.fields] + [r["_rows"]]


def filename(result: Result, fmt: str) -> str:
    release = result.release.number if result.release else 0
    by = "_".join(result.query.by) or "total"
    return f"cedars_{result.query.dataset}_release{release}_{by}.{fmt}"


def render(result: Result, fmt: str) -> bytes:
    if fmt not in FORMATS:
        raise ValueError(f"unknown format {fmt!r}")
    header, rows = table(result, totals=fmt != "parquet")
    stamp = result.stamp() | {"generated_at": timezone.now().isoformat(), "notes": result.notes}
    return {"csv": _csv, "xlsx": _xlsx, "parquet": _parquet}[fmt](header, rows, stamp)


def _stamp_lines(stamp: dict) -> list[tuple[str, str]]:
    return [(k, v if isinstance(v, str) else json.dumps(v, default=str)) for k, v in stamp.items()]


def _csv(header, rows, stamp) -> bytes:
    out = io.StringIO()
    for k, v in _stamp_lines(stamp):
        out.write(f"# {k}: {v}\n")
    w = csv.writer(out, lineterminator="\n")
    w.writerow(header)
    w.writerows(rows)
    return out.getvalue().encode("utf-8")


def _xlsx(header, rows, stamp) -> bytes:
    from openpyxl import Workbook

    wb = Workbook(write_only=True)
    data = wb.create_sheet("Data")
    data.append(header)
    for r in rows:
        data.append(r)
    about = wb.create_sheet("About")
    about.append(["Item", "Value"])
    for k, v in _stamp_lines(stamp):
        about.append([k, v])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _parquet(header, rows, stamp) -> bytes:
    import pyarrow as pa
    import pyarrow.parquet as pq

    columns = list(zip(*rows, strict=True)) if rows else [[] for _ in header]
    names = _unique(header)
    t = pa.table({n: list(c) for n, c in zip(names, columns, strict=True)})
    t = t.replace_schema_metadata({"cedars": json.dumps(stamp, default=str)})
    out = io.BytesIO()
    pq.write_table(t, out)
    return out.getvalue()


def _unique(names: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for n in names:
        seen[n] = seen.get(n, 0) + 1
        out.append(n if seen[n] == 1 else f"{n} ({seen[n]})")
    return out
