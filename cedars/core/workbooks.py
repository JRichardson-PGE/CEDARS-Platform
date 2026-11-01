"""Every workbook the platform writes (BRD v0.5 BR-10.3: published documents meet WCAG 2.2 AA): the budget filing appendix, the record
and dashboard downloads, and the upload templates. One way to write them, so each is accessible the same way:

- the workbook has a title, a subject and a language (en-US) in its properties;
- every sheet has a name that says what it holds (never Excel's "Sheet1");
- every sheet's first row is its header, marked as such: the rows are an Excel table with a header row (what screen readers announce
  column by column, and what Excel's own accessibility checker asks for), the row is repeated on printed pages (print titles) and
  frozen on screen, and its names are unique;
- no merged cells, pictures or charts: a layout is rows under one header, and anything shown as a picture on a page is a table here.

The document accessibility check (cedars.services.accessibility.checks) finds nothing in a workbook written this way; the tests hold
every generated workbook to that. Write-only workbooks (large downloads) stream their rows: `Book(write_only=True)`."""

from __future__ import annotations

import io
import re
import warnings

LANGUAGE = "en-US"
MIN_WIDTH, MAX_WIDTH = 10, 60


def unique_header(names) -> list[str]:
    """The header as a table needs it: text, none blank, no two alike (compared without case, as Excel does)."""
    out, seen = [], set()
    for i, raw in enumerate(names, start=1):
        name = str(raw if raw is not None else "").strip() or f"Column {i}"
        candidate, n = name, 2
        while candidate.lower() in seen:
            candidate, n = f"{name} ({n})", n + 1
        seen.add(candidate.lower())
        out.append(candidate)
    return out


def _table_name(sheet: str, taken: set[str]) -> str:
    base = re.sub(r"[^A-Za-z0-9_]", "", sheet.title().replace(" ", "")) or "Data"
    if not base[0].isalpha():
        base = f"T{base}"
    name, n = f"{base}Table", 2
    while name.lower() in taken:
        name, n = f"{base}Table{n}", n + 1
    taken.add(name.lower())
    return name


class Sheet:
    """One sheet: its header is written when it is made, rows are appended, and `finish` (Book.save calls it) marks the rows as a table."""

    def __init__(self, book: Book, name: str, header, *, widths: dict[int, int] | None = None):
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter

        self.book, self.name = book, name
        self.header = unique_header(header)
        self.rows = 0
        self.ws = book.wb.create_sheet(name)
        self.ws.freeze_panes = "A2"
        for i, h in enumerate(self.header, start=1):
            width = (widths or {}).get(i - 1) or min(MAX_WIDTH, max(MIN_WIDTH, len(h) + 4))
            self.ws.column_dimensions[get_column_letter(i)].width = width
        if book.write_only:
            from openpyxl.cell import WriteOnlyCell

            cells = []
            for h in self.header:
                c = WriteOnlyCell(self.ws, value=h)
                c.font = Font(bold=True)
                cells.append(c)
            self.ws.append(cells)
        else:
            self.ws.append(self.header)
            for c in self.ws[1]:
                c.font = Font(bold=True)
        self.finished = False

    def append(self, row) -> None:
        self.ws.append(list(row))
        self.rows += 1

    def extend(self, rows) -> None:
        for r in rows:
            self.append(r)

    def finish(self) -> None:
        if self.finished:
            return
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.filters import AutoFilter
        from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo

        # a table needs at least one row under its header (an empty one stays empty, ready to fill in)
        ref = f"A1:{get_column_letter(len(self.header))}{max(self.rows, 1) + 1}"
        t = Table(displayName=_table_name(self.name, self.book.table_names), ref=ref)
        t.tableColumns = [TableColumn(id=i, name=h) for i, h in enumerate(self.header, start=1)]
        t.autoFilter = AutoFilter(ref=ref)
        # Light 1: black text, no colored fills, so the cells keep their contrast in any theme
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", "In write-only mode you must add table columns manually")
            self.ws.add_table(t)
        self.ws.print_title_rows = "1:1"
        self.finished = True


class Book:
    def __init__(self, title: str, *, subject: str = "", description: str = "", write_only: bool = False):
        from openpyxl import Workbook

        self.write_only = write_only
        self.wb = Workbook(write_only=write_only)
        if not write_only:
            self.wb.remove(self.wb.active)  # every sheet is made with its name and header
        p = self.wb.properties
        p.title, p.subject, p.description = title, subject or title, description
        p.creator = "CEDARS"
        p.language = LANGUAGE
        self.sheets: list[Sheet] = []
        self.table_names: set[str] = set()

    def sheet(self, name: str, header, rows=(), **kw) -> Sheet:
        if len(name) > 31:
            raise ValueError(f"A sheet name has at most 31 characters: {name}")
        s = Sheet(self, name, header, **kw)
        s.extend(rows)
        self.sheets.append(s)
        return s

    def save(self, out=None) -> bytes | None:
        """Write the workbook to `out` (a path or binary file); with no `out`, return its bytes."""
        for s in self.sheets:
            s.finish()
        if out is not None:
            self.wb.save(out)
            return None
        buf = io.BytesIO()
        self.wb.save(buf)
        return buf.getvalue()
