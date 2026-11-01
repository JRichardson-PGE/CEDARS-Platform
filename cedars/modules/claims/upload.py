"""A claims upload: a .zip of the template's tables as CSV files (comma-, pipe- or tab-delimited, named after the table: Claim.csv,
Measure.csv, ...) or an Excel workbook with a sheet per table (BR-02.2). A quarter's upload holds the records that are new or changed
since the last confirmed quarter, so any of the tables may be left out. Columns are read under their canonical names (CEDARS' names
and CET's are accepted, decision L); a column the template does not have is ignored, not rejected (BR-02.1).

The PII tables never come this way: a file named PII_Contact or PII_Site refuses the whole upload (BR-10.1, BR-10.4)."""

import csv
import io
import zipfile
from dataclasses import dataclass, field
from pathlib import PurePosixPath

from cedars.core import uploads

from . import spec

SKIPPED = ("__MACOSX/", ".DS_Store", "Thumbs.db", "desktop.ini")


class UploadError(ValueError):
    """The upload is not a claims upload; the message says why, in words for the person who uploaded it."""


@dataclass
class Table:
    """One table of an upload, as text: rows keyed by canonical column names, in file order (row i is line i + 2 of the file)."""

    name: str
    columns: list[str]
    rows: list[dict[str, str]]
    ignored: list[str] = field(default_factory=list)

    def get(self, row: dict, column: str) -> str:
        return row.get(column, "")


def combine(files: list[tuple[str, bytes]]) -> uploads.Combined:
    """Several files chosen at once as the one upload read takes (cedars.core.uploads, BR-06.3); a PII table in any of them refuses
    the whole upload, as in one file."""
    try:
        return uploads.combine(files, spec.table_name, refuse=_no_pii)
    except uploads.CombineError as e:
        raise UploadError(str(e)) from None


def read(data: bytes, filename: str) -> dict[str, Table]:
    """{table name: Table} from an uploaded .zip or .xlsx."""
    raw = _from_workbook(data) if filename.lower().endswith(".xlsx") else _from_zip(data)
    if not raw:
        raise UploadError(f"The file holds none of the claims tables ({', '.join(spec.TABLES)}).")
    return {name: _table(name, content) for name, content in raw.items()}


def _from_zip(data: bytes) -> dict[str, list[list[str]]]:
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise UploadError(
            "The file is not a valid .zip or .xlsx. Upload a .zip of the claims tables as CSV files, or an Excel workbook with a sheet "
            "per table."
        ) from None
    out, extra = {}, []
    for info in z.infolist():
        name = info.filename
        if info.is_dir() or any(s in name for s in SKIPPED) or PurePosixPath(name).name.startswith("._"):
            continue
        p = PurePosixPath(name)
        if len(p.parts) > 2 or p.suffix.lower() not in (".csv", ".txt"):
            extra.append(name)
            continue
        _no_pii(p.stem)
        table = spec.table_name(p.stem)
        if table is None:
            extra.append(name)
        elif table in out:
            raise UploadError(f"The .zip holds two {table} files; keep one.")
        else:
            out[table] = _csv_rows(z.read(info))
    if extra:
        raise UploadError(f"The .zip may hold only the claims tables ({', '.join(spec.TABLES)}); it also has {', '.join(sorted(extra))}.")
    return out


def _no_pii(name: str) -> None:
    if spec.is_pii_table(name):
        raise UploadError(
            f"{name} holds personal information, which never travels with the claim records. Remove it from this file and send it "
            "through the PII upload; the claim records carry only the SiteID and ContactID keys."
        )


def _from_workbook(data: bytes) -> dict[str, list[list[str]]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # openpyxl raises several kinds for a file that is not a workbook
        raise UploadError("The file is not a valid Excel workbook (.xlsx).") from None
    out = {}
    for ws in wb.worksheets:
        _no_pii(ws.title)
        table = spec.table_name(ws.title)
        if table is None:
            continue
        rows = []
        for row in ws.iter_rows(values_only=True):
            if row is None or all(c is None or str(c).strip() == "" for c in row):
                continue
            rows.append(["" if c is None else _cell(c) for c in row])
        out[table] = rows
    wb.close()
    return out


def _cell(value) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if hasattr(value, "date") and hasattr(value, "hour"):  # a datetime from a date cell
        return value.date().isoformat() if not (value.hour or value.minute or value.second) else value.isoformat()
    return str(value)


def text(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        return content.decode("cp1252")


def _csv_rows(content: bytes) -> list[list[str]]:
    body = text(content)
    first = body.split("\n", 1)[0]
    sep = "|" if "|" in first else ("\t" if "\t" in first else ",")
    return [r for r in csv.reader(io.StringIO(body), delimiter=sep) if r and any(c.strip() for c in r)]


def _table(name: str, raw: list[list[str]]) -> Table:
    if not raw:
        return Table(name, [], [])
    header = [h.strip() for h in raw[0]]
    names = [spec.canonical(name, h) for h in header]
    seen: set[str] = set()
    for h, n in zip(header, names, strict=True):
        if n is not None and n in seen:
            raise UploadError(f"{name} has the column {n} twice (as {h} and another name); keep one.")
        if n is not None:
            seen.add(n)
    rows = []
    for r in raw[1:]:
        rows.append({n: (r[i].strip() if i < len(r) else "") for i, n in enumerate(names) if n is not None})
    columns = [n for n in names if n is not None]
    return Table(name, columns, rows, ignored=[h for h, n in zip(header, names, strict=True) if n is None])


def table_from_csv(name: str, content: bytes) -> Table:
    """One table's CSV read as an upload reads it (the pre-upload checker)."""
    return _table(name, _csv_rows(content))


def without_system_fields(t: Table) -> Table:
    """An upload's table without the columns CEDARS fills in (spec.SYSTEM_FIELDS): they are ignored, as an unknown column is."""
    drop = [c for c in spec.SYSTEM_FIELDS.get(t.name, ()) if c in t.columns]
    if not drop:
        return t
    rows = [{k: val for k, val in r.items() if k not in drop} for r in t.rows]
    return Table(t.name, [c for c in t.columns if c not in drop], rows, ignored=t.ignored + drop)


# --- writing tables


def to_csv(columns: list[str], rows: list[dict[str, str]]) -> bytes:
    out = io.StringIO()
    w = csv.writer(out, lineterminator="\n")
    w.writerow(columns)
    for r in rows:
        w.writerow([r.get(c, "") for c in columns])
    return out.getvalue().encode("utf-8")


def to_zip(tables: dict[str, Table]) -> bytes:
    """Tables as a .zip of comma-delimited CSVs in the template's order (an empty table is left out)."""
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for name in spec.TABLES:
            t = tables.get(name)
            if t is not None and t.rows:
                z.writestr(f"{name}.csv", to_csv(t.columns, t.rows))
    return out.getvalue()


def from_zip(data: bytes) -> dict[str, Table]:
    """Read back a .zip this module wrote (to_zip)."""
    raw = _from_zip(data) if data else {}
    return {name: _table(name, content) for name, content in raw.items()}
