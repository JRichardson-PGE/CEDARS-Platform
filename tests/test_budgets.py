import io
import json
import zipfile
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path

import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from django.urls import reverse
from django.utils import timezone

from cedars.core import storage
from cedars.modules.budgets import jobfile, service
from cedars.modules.budgets.models import BudgetLine
from cedars.modules.calculations import service as calculations
from cedars.modules.calculations.models import Run
from cedars.services.accounts import service as access
from cedars.services.accounts.service import AccessDenied
from cedars.services.audit.models import AuditEvent
from cedars.services.registry import service as registry
from cedars.services.submissions import service as submissions
from cedars.services.submissions.models import SubmissionVersion

from .conftest import sign_in

pytestmark = pytest.mark.django_db
DATA = Path(__file__).parent / "data"
COSTS = "PrgID|PrgYear|ClaimYearQuarter|PA|AdminCostsOverheadAndGA|AdminCostsOther|MarketingOutreach|DIActivity|DIInstallation|DIHardwareAndMaterials|DIRebateAndInspection|EMV|UserInputIncentive|OnBillFinancing|CostsRecoveredFromOtherSources"  # noqa: E501
MEASURE = "CEInputID|PrgID|ClaimYearQuarter|NumUnits|UnitEndUserRebate|UnitIncentiveToOthers|UnitDirectInstallLab|UnitDirectInstallMat"


def job(costs=None, measures=None, extra=None) -> bytes:
    costs = costs or [
        "PGE_Res_HVAC|2028|2028Q1|PGE|1000|500|200|0|0|0|0|100|0|5000|0",
        "PGE_Res_HVAC|2029|2029Q1|PGE|1100|500|200|0|0|0|0|100|0|0|0",
        "PGE_Com_Ltg|2028|2028Q1|PGE|300|0|0|0|0|0|0|0|0|0|0",
    ]
    measures = measures or [
        "M1|PGE_Res_HVAC|2028Q1|10|20|30|5|5",
        "M2|PGE_Res_HVAC|2029Q2|2|100|0|0|0",
        "M3|PGE_Com_Ltg|2028Q3|4|25|0|0|0",
    ]
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("ProgramCost.csv", "\n".join([COSTS] + costs) + "\n")
        z.writestr("Measure.csv", "\n".join([MEASURE] + measures) + "\n")
        for k, v in (extra or {}).items():
            z.writestr(k, v)
    return out.getvalue()


@pytest.fixture
def programs(pge_admin, pge, cpuc_admin):
    sce = access.create_organisation(cpuc_admin, code="SCE", name="Southern California Edison", kind="pa", pa_type="IOU")
    registry.create_program(pge_admin, prg_id="PGE_Res_HVAC", administrator=pge, valid_from=date(2024, 1, 1), name="Residential HVAC")
    registry.create_program(pge_admin, prg_id="PGE_Com_Ltg", administrator=pge, valid_from=date(2024, 1, 1), name="Commercial lighting")
    registry.create_program(cpuc_admin, prg_id="SCE_Ind", administrator=sce, valid_from=date(2024, 1, 1), name="Industrial")
    return sce


@pytest.fixture
def filer(make_user, pge, pge_admin):
    u = make_user("budgets@pge.com", "Bea Udget")
    for p in ("budgets.submit", "budgets.confirm"):
        access.grant(pge_admin, u, pge, p)
    return u


@pytest.fixture
def filing(cpuc_admin):
    return service.add_filing(
        cpuc_admin, first_year=2028, kind="True-up advice letter (TUAL)", due_at=timezone.now() + timedelta(days=30), acc_version="2026"
    )


@pytest.fixture
def release(cpuc_admin):
    return calculations.add_release(cpuc_admin, release="0.1.80", calculations=["recet"], accs={"recet": ["2024", "2026"]}, default=True)


def group_sums(scale: float = 1.0) -> bytes:
    """A run's kept group sums (out/_results/group_sums2_*.parquet), as ReCET's results_view writes them."""
    rows = [
        {"_pa": "PGE", "_prg": "PGE_Res_HVAC", "_prgname": "Res HVAC", "_sector": "Res", "_segment": "RA", "_year": "2028", "_n": 1,
         "TotalSystemBenefit": 3000.0 * scale, "TRCBen": 3000.0 * scale, "TRCCost": 2000.0, "PACBen": 3000.0 * scale, "PACCost": 1500.0,
         "FirstYearNetKWh": 100.0},
        {"_pa": "PGE", "_prg": "PGE_Res_HVAC", "_prgname": "Res HVAC", "_sector": "Com", "_segment": "RA", "_year": "2028", "_n": 1,
         "TotalSystemBenefit": 1000.0, "TRCBen": 1000.0, "TRCCost": 1000.0, "PACBen": 1000.0, "PACCost": 500.0, "FirstYearNetKWh": 50.0},
        {"_pa": "PGE", "_prg": "PGE_Res_HVAC", "_prgname": "Res HVAC", "_sector": "Res", "_segment": "RA", "_year": "2029", "_n": 1,
         "TotalSystemBenefit": 500.0, "TRCBen": 500.0, "TRCCost": 1000.0, "PACBen": 500.0, "PACCost": 400.0, "FirstYearNetKWh": 20.0},
        {"_pa": "PGE", "_prg": "PGE_Com_Ltg", "_prgname": "Com Ltg", "_sector": "Com", "_segment": "RA", "_year": "2028", "_n": 1,
         "TotalSystemBenefit": 400.0, "TRCBen": 400.0, "TRCCost": 100.0, "PACBen": 400.0, "PACCost": 100.0, "FirstYearNetKWh": None},
    ]  # fmt: skip
    buf = io.BytesIO()
    pq.write_table(pa.Table.from_pylist(rows), buf)
    return buf.getvalue()


def finish_run(run: Run, scale: float = 1.0) -> Run:
    """The worker's end, as a ReCET worker leaves it: group sums under out/_results/ and a receipt."""
    calculations.mark_started(run, worker="test")
    out = run.storage_prefix + "out/"
    storage.store("results").put_bytes(out + "_results/group_sums2_measure_level_run.parquet", group_sums(scale))
    storage.store("results").put_bytes(out + "receipt.json", b"{}")
    return calculations.finish(run, exit_code=0, receipt={"packs": {"files": {}}, "totals": {"TRCRatio": 1.5}})


def priced_version(filer, filing, pge, release):
    v = service.upload(filer, filing, pge, job(), "tual.zip")
    run = Run.objects.get(pk=service.price(filer, v))
    finish_run(run)
    v.refresh_from_db()
    return v, run


# --- the upload


def test_a_zip_or_a_workbook_is_a_job():
    tables, data = jobfile.read_upload(job(extra={"__MACOSX/._Measure.csv": "x", "jobs/.DS_Store": "x"}), "f.zip")
    assert set(tables) == {"Measure", "ProgramCost"} and data[:2] == b"PK"
    with pytest.raises(jobfile.JobError, match="only Measure.csv and ProgramCost.csv; it also has notes.docx"):
        jobfile.read_upload(job(extra={"notes.docx": "x"}), "f.zip")
    with pytest.raises(jobfile.JobError, match="not a valid .zip"):
        jobfile.read_upload(b"PrgID,PrgYear\n", "f.zip")
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "ProgramCost"
    for line in [COSTS, "PGE_Res_HVAC|2028|2028Q1|PGE|1000|0|0|0|0|0|0|0|0|0|0"]:
        ws.append([float(c) if c.isdigit() else c for c in line.split("|")])
    m = wb.create_sheet("measure")
    m.append(MEASURE.split("|"))
    m.append(["M1", "PGE_Res_HVAC", "2028Q1", 10, 20, 0, 0, 0])
    buf = io.BytesIO()
    wb.save(buf)
    tables, data = jobfile.read_upload(buf.getvalue(), "tual.xlsx")
    assert zipfile.ZipFile(io.BytesIO(data)).namelist() == ["Measure.csv", "ProgramCost.csv"]
    header, rows = jobfile.rows(tables["ProgramCost"])
    assert rows[0]["PrgYear"] == "2028" and rows[0]["AdminCostsOverheadAndGA"] == "1000"
    del wb["measure"]
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(jobfile.JobError, match="no Measure sheet"):
        jobfile.read_upload(buf.getvalue(), "tual.xlsx")


def test_only_a_cpuc_system_administrator_sets_up_filings(cpuc_admin, pge_admin, filing):
    with pytest.raises(AccessDenied, match="system administrator"):
        service.add_filing(pge_admin, first_year=2028, kind="MCAL", due_at=timezone.now(), acc_version="2026")
    with pytest.raises(AccessDenied, match="already has"):
        service.add_filing(cpuc_admin, first_year=2028, kind=filing.kind, due_at=timezone.now(), acc_version="2026")
    draft = registry.import_allocation_csv(cpuc_admin, (DATA / "statewide_factors_20260918.csv").read_bytes(), label="d", purpose="filing")
    with pytest.raises(AccessDenied, match="draft"):
        service.add_filing(cpuc_admin, first_year=2028, kind="MCAL", due_at=timezone.now(), acc_version="2026", allocation_version=draft)
    assert filing.cycle.first_year == 2028 and filing.cycle.last_year == 2031 and filing.budget_setup.acc_version == "2026"
    assert AuditEvent.objects.filter(action="budget_filing.set_up").exists()


def test_an_upload_gives_the_budget_by_program_and_year(filer, filing, pge, programs):
    v = service.upload(filer, filing, pge, job(), "tual.zip")
    assert v.status == SubmissionVersion.Status.VALIDATED and v.findings["records"] == 6 and v.findings["rejected"] == 0
    lines = {(ln.prg_id, ln.year): ln for ln in BudgetLine.objects.filter(version=v)}
    hvac = lines[("PGE_Res_HVAC", 2028)]
    assert sum(Decimal(x) for x in hvac.program_costs.values()) == 1800  # the ten ProgramCost fields CET adds up
    assert hvac.rebates_and_incentives == 10 * (20 + 30 + 5 + 5) and hvac.measures == 1 and hvac.on_bill_financing == 5000
    assert lines[("PGE_Res_HVAC", 2029)].rebates_and_incentives == 200
    t = service.table(v)
    assert t.total.budget == 1800 + 600 + 1900 + 200 + 300 + 100 and not t.priced
    assert [r.prg_id for r in service.table(v, years=["2029"]).rows] == ["PGE_Res_HVAC"]


def test_module_checks_reject_programs_and_years_outside_the_filing(filer, filing, pge, programs):
    bad = job(
        costs=["PGE_Res_HVAC|2027|2027Q1|PGE|1|0|0|0|0|0|0|0|0|0|0", "SCE_Ind|2028|2028Q1|SCE|1|0|0|0|0|0|0|0|0|0|0",
               "PGE_Com_Ltg|2028|2028Q1|PGE|ten|0|0|0|0|0|0|0|0|0|0"],
        measures=["M1|NOPE|2028Q1|1|0|0|0|0", "M2|PGE_Res_HVAC|2029Q1|1|0|0|0|0"],
    )  # fmt: skip
    v = service.upload(filer, filing, pge, bad, "tual.zip")
    assert v.status == SubmissionVersion.Status.DRAFT
    found = {(f["table"], f["row"], f["rule_id"], f["severity"]) for f in service.findings(v)}
    assert ("ProgramCost", 2, "BUD-YEAR", "reject") in found and ("ProgramCost", 3, "BUD-ORG", "reject") in found
    assert ("ProgramCost", 4, "BUD-NUM", "reject") in found and ("Measure", 2, "BUD-PRG", "reject") in found
    assert ("Measure", 0, "BUD-NOCOST", "warning") in found  # M2: its program has no 2029 ProgramCost row
    assert v.findings["rejected"] == 4 and v.findings["by_rule"][0]["severity"] == "reject"
    assert "BUD-YEAR" in service.findings_csv(v)
    with pytest.raises(AccessDenied, match="validated"):
        service.price(filer, v)


def test_a_new_upload_replaces_the_last_and_keeps_nothing_superseded(filer, filing, pge, programs, django_capture_on_commit_callbacks):
    v = service.upload(filer, filing, pge, job(), "tual.zip")
    first = service.job_file(v)["key"]
    with django_capture_on_commit_callbacks(execute=True):
        v = service.upload(filer, filing, pge, job(costs=["PGE_Com_Ltg|2028|2028Q1|PGE|5|0|0|0|0|0|0|0|0|0|0"]), "tual2.zip")
    assert v.number == 1 and not storage.store("uploads").exists(first)
    assert set(BudgetLine.objects.filter(version=v).values_list("prg_id", "year")) == {
        ("PGE_Com_Ltg", 2028),
        ("PGE_Res_HVAC", 2028),
        ("PGE_Res_HVAC", 2029),
    }


def test_a_filer_needs_budgets_submit(make_user, pge, pge_admin, filing, programs):
    u = make_user("reader@pge.com")
    access.grant(pge_admin, u, pge, "budgets.read")
    with pytest.raises(AccessDenied, match="budgets.submit"):
        service.upload(u, filing, pge, job(), "tual.zip")


# --- priced, confirmed, published


def test_priced_then_confirmed_and_published(filer, filing, pge, programs, release, make_user):
    v, run = priced_version(filer, filing, pge, release)
    assert run.submission_version_id == v.pk and run.acc_version == "2026" and run.settings["allocate_statewide"] is False
    assert service.priced(v) and v.calculation["job_sha256"] == service.job_file(v)["sha256"]
    assert v.calculation["results"]["group_sums"].endswith("group_sums2_measure_level_run.parquet")
    t = service.table(v)
    hvac = next(r for r in t.rows if r.prg_id == "PGE_Res_HVAC")
    assert hvac.value("TotalSystemBenefit") == 4500 and hvac.value("TRC") == pytest.approx(4500 / 4000)
    assert service.table(v, years=["2028"]).total.value("TRC") == pytest.approx(4400 / 3100)
    community = make_user("someone@example.com")
    assert not service.can_read(community, v)
    v = service.confirm(filer, v)
    assert v.status == SubmissionVersion.Status.PUBLISHED and service.can_read(community, v)
    assert [label for label, *_ in service.published_files(v)] == [
        "Inputs (Measure and ProgramCost)",
        "Run settings (job spec)",
        "Run receipt (engine, reference data, outputs)",
    ]
    assert service.standing(pge, filing.cycle) == v


def test_a_changed_upload_must_be_priced_again(filer, filing, pge, programs, release):
    v, _ = priced_version(filer, filing, pge, release)
    v = service.upload(filer, filing, pge, job(costs=["PGE_Com_Ltg|2028|2028Q1|PGE|5|0|0|0|0|0|0|0|0|0|0"]), "tual.zip")
    assert not service.priced(v) and v.calculation == {}
    with pytest.raises(AccessDenied, match="Price this version first"):
        service.confirm(filer, v)


def test_excluded_programs_are_filtered_never_zeroed(filer, filing, pge, cpuc_admin, programs, release):
    v, _ = priced_version(filer, filing, pge, release)
    ltg = registry.catalog()
    ltg = next(p for p, _ in ltg if p.prg_id == "PGE_Com_Ltg")
    registry.change_program(cpuc_admin, ltg, effective=date(2025, 1, 1), reason="loan pool", exclude_from_budget=True)
    t = service.table(v)
    row = next(r for r in t.rows if r.prg_id == "PGE_Com_Ltg")
    assert row.excluded_budget and row.budget == 400
    hidden = service.table(v, hide_budget_excluded=True)
    assert "PGE_Com_Ltg" not in [r.prg_id for r in hidden.rows] and hidden.total.budget == t.total.budget - 400
    assert t.reconciliation["filed"] == t.total.budget and t.reconciliation["without_excluded"] == t.total.budget - 400
    assert t.reconciliation["on_bill_financing"] == 5000


def test_a_filing_pins_its_allocation_version(filer, filing, pge, programs, release, cpuc_admin, pge_admin):
    factors = (DATA / "statewide_factors_20260918.csv").read_bytes()
    first = registry.import_allocation_csv(cpuc_admin, factors, label="2028 TUAL", purpose="filing")
    first = registry.confirm_allocation(cpuc_admin, first)
    service.change_setup(cpuc_admin, filing.budget_setup, acc_version="2026", allocation_version=first, reason="the TUAL's factors")
    sw = factors.split(b"\n")[1].split(b",")[1].decode()
    registry.create_program(cpuc_admin, prg_id=sw, administrator=None, statewide=True, valid_from=date(2020, 1, 1), name="SW")
    v = service.upload(
        filer, filing, pge, job(costs=[f"{sw}|2028|2028Q1|PGE|10|0|0|0|0|0|0|0|0|0|0"], measures=[f"M1|{sw}|2028Q1|1|0|0|0|0"]), "t.zip"
    )
    run = Run.objects.get(pk=service.price(filer, v))
    finish_run(run)
    v.refresh_from_db()
    assert run.allocation_version == first and v.calculation["allocation_version"]["label"] == "2028 TUAL"
    later = registry.import_allocation_csv(cpuc_admin, factors, label="2028 claims", purpose="claims")
    registry.confirm_allocation(cpuc_admin, later)
    v.refresh_from_db()
    assert v.calculation["allocation_version"]["label"] == "2028 TUAL"  # BR-04.8: a later version never changes it


def test_two_versions_compare_by_program(filer, filing, pge, programs, release, cpuc_admin):
    a, _ = priced_version(filer, filing, pge, release)
    a = service.confirm(filer, a)
    mcal = service.add_filing(cpuc_admin, first_year=2028, kind="MCAL", due_at=timezone.now() + timedelta(days=60), acc_version="2026")
    b = service.upload(filer, mcal, pge, job(costs=["PGE_Res_HVAC|2028|2028Q1|PGE|2000|0|0|0|0|0|0|0|0|0|0"]), "mcal.zip")
    rows = {r["prg_id"]: r for r in service.compare(a, b)}
    budget_a, budget_b, diff = rows["PGE_Res_HVAC"]["cells"][0]
    assert (budget_a, budget_b, diff) == (Decimal(2400 + 2100), Decimal(2000 + 600 + 200), 2800 - 4500)
    assert rows["Total"]["total"] and len(service.compare(a, a, changed_only=True)) == 1  # only the total


def test_the_appendix_is_built_from_the_version(filer, filing, pge, programs, release):
    from openpyxl import load_workbook

    v, _ = priced_version(filer, filing, pge, release)
    wb = load_workbook(io.BytesIO(service.appendix_xlsx(v)))
    assert wb.sheetnames == ["About", "Budget", "Cost-effectiveness", "Savings by year"]
    budget = list(wb["Budget"].iter_rows(values_only=True))
    assert budget[0][:3] == ("PrgID", "Year", "PA") and budget[1][0] == "PGE_Com_Ltg"
    assert dict((r[0], r[1]) for r in wb["About"].iter_rows(values_only=True))["ACC"] == "2026 ACC"


def test_after_the_due_date_a_correction_is_a_new_version(
    filer, filing, pge, programs, release, cpuc_admin, django_capture_on_commit_callbacks
):
    v, _ = priced_version(filer, filing, pge, release)
    v = service.confirm(filer, v)
    submissions.move_deadline(cpuc_admin, filing, timezone.now() - timedelta(minutes=1), "closed")
    submissions.set_policy(cpuc_admin, "budgets", reopen_after_due="open")
    v2 = submissions.request_reopen(filer, v.submission, "one program's 2029 costs were wrong")
    assert v2.number == 2 and v2.files == v.files
    service.ensure_lines(v2)
    assert v2.budget_lines.count() == v.budget_lines.count()
    assert SubmissionVersion.objects.get(pk=v.pk).status == SubmissionVersion.Status.PUBLISHED
    with django_capture_on_commit_callbacks(execute=True):  # a correction's upload leaves the published version's files alone
        service.upload(filer, filing, pge, job(), "correction.zip")
    assert all(storage.store("uploads").exists(f["key"]) for f in v.files)


# --- pages


def test_the_pages(client, filer, filing, pge, programs, release, cpuc_admin, make_user):
    sign_in(client, filer)
    assert client.get(reverse("budgets:index")).status_code == 200
    page = client.get(reverse("budgets:filing", args=[filing.pk]))
    assert page.status_code == 200 and b"Upload the filing" in page.content
    from django.core.files.uploadedfile import SimpleUploadedFile

    r = client.post(reverse("budgets:upload", args=[filing.pk]), {"file": SimpleUploadedFile("tual.zip", job())})
    assert r.status_code == 302
    v = SubmissionVersion.objects.get()
    assert v.status == "validated"
    assert b"Price this version" in client.get(reverse("budgets:filing", args=[filing.pk])).content
    client.post(reverse("budgets:price", args=[filing.pk]))
    run = Run.objects.get()
    assert b"the run's page" in client.get(reverse("budgets:filing", args=[filing.pk])).content
    finish_run(run)
    assert b"Confirm and publish" in client.get(reverse("budgets:filing", args=[filing.pk])).content
    client.post(reverse("budgets:confirm", args=[filing.pk]))
    v.refresh_from_db()
    assert v.status == "published"
    page = client.get(reverse("budgets:version", args=[v.pk]) + "?years=2028&hide_budget=1")
    assert page.status_code == 200 and b"PGE_Res_HVAC" in page.content
    csv = client.get(reverse("budgets:version_csv", args=[v.pk]) + "?years=2029").content.decode()
    assert csv.splitlines()[1].startswith("PGE_Res_HVAC") and len(csv.splitlines()) == 3
    assert client.get(reverse("budgets:file", args=[v.pk, 0])).status_code == 200
    assert client.get(reverse("budgets:appendix", args=[v.pk])).status_code == 200
    assert client.get(reverse("budgets:compare") + f"?a={v.pk}&b={v.pk}").status_code == 200
    # a community user sees the published version but not the findings or another PA's draft
    community = make_user("someone@example.com")
    client.logout()
    sign_in(client, community)
    assert client.get(reverse("budgets:version", args=[v.pk])).status_code == 200
    assert client.get(reverse("budgets:findings_csv", args=[v.pk])).status_code == 404
    assert client.get(reverse("budgets:filing", args=[filing.pk]) + "?org=PGE").status_code == 404
    assert client.get(reverse("budgets:filing_new")).status_code == 404
    client.logout()
    sign_in(client, cpuc_admin)
    assert client.get(reverse("budgets:filing", args=[filing.pk]) + "?org=PGE").status_code == 200
    r = client.post(
        reverse("budgets:filing_new"),
        {"first_year": 2028, "kind": "Mid-cycle advice letter (MCAL)", "due": "2029-09-01", "acc_version": "2026"},
    )
    assert r.status_code == 302
    assert client.get(reverse("budgets:setup", args=[filing.pk])).status_code == 200


def test_the_results_file_names_the_settings_it_was_priced_with(filer, filing, pge, programs, release):
    v, run = priced_version(filer, filing, pge, release)
    spec = json.loads(storage.store("results").get_bytes(v.calculation["spec_key"]))
    assert spec["settings"]["acc_version"] == "2026" and spec["inputs"]["job_zip"]["sha256"] == v.calculation["job_sha256"]


def test_the_budget_filing_rules_run_year_by_year(filer, filing, pge, programs, cpuc_admin):
    pytest.importorskip("recet.filing_rules", reason="the rule library runs on the pinned ReCET engine")
    from cedars.services.rules import service as rules

    rs = rules.import_from_engine(cpuc_admin, "budget_filing")
    rules.publish(cpuc_admin, rs)
    v = service.upload(filer, filing, pge, job(), "tual.zip")  # a file without the specification's columns
    assert {f["rule_id"] for f in service.findings(v)} == {"BUD-COL"} and v.status == "draft"
    assert any("were not run" in n for n in v.findings["notes"])
    tables = {}
    for table, rows in (("ProgramCost", job_rows(COSTS)), ("Measure", job_rows(MEASURE))):
        spec = sorted(set(rs.rules.filter(table=table, kind="spec").values_list("field", flat=True)))
        cols = list(dict.fromkeys(list(rows[0]) + spec))
        tables[table] = "\n".join(["|".join(cols)] + ["|".join(r.get(c, "") for c in cols) for r in rows]) + "\n"
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for k, text in tables.items():
            z.writestr(f"{k}.csv", text)
    v = service.upload(filer, filing, pge, out.getvalue(), "full.zip")
    found = service.findings(v)
    assert v.findings["rule_set"].startswith(rs.label) and found and "BUD-COL" not in {f["rule_id"] for f in found}
    # a cycle filing spans four years: CEDARS' single-year rules run per year, so no 2029 row is rejected for not being 2028
    assert not [f for f in found if f["table"] == "ProgramCost" and f["field"] == "PrgYear"]
    assert {f["row"] for f in found if f["table"] == "ProgramCost"} == {2, 3, 4}  # the file's own line numbers


def job_rows(header: str) -> list[dict]:
    tables, _ = jobfile.read_upload(job(), "j.zip")
    name = "ProgramCost" if header == COSTS else "Measure"
    return jobfile.rows(tables[name])[1]
