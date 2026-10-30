"""A monthly report's upload: a .zip holding Commitments.csv, Measure.csv and ProgramCost.csv (comma-, pipe- or tab-delimited), or an
Excel workbook with those three sheets (BR-02.2). Measure and ProgramCost are the year-to-date BCA inputs: the month's spend is read from
them (Jake, 2026-10-30), and they become the version's job .zip, two CSVs as a calculation run takes them, priced for the estimates
(BR-14.5). Commitments is its own table (D2)."""

import csv
import hashlib
import io
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import PurePosixPath

from .models import COMMITMENT_FIELDS

BCA_TABLES = ("Measure", "ProgramCost")
TABLES = ("Commitments",) + BCA_TABLES
REQUIRED = {
    "Commitments": ["PA", "PrgID", "YearMonth"] + [f for f, _ in COMMITMENT_FIELDS],
    "Measure": ["CEInputID", "PrgID", "ClaimYearQuarter", "NumUnits"],
    "ProgramCost": ["PrgID", "PrgYear", "PA"],
}
# Columns of CEDARS' monthly Program table that Commitments does not read, and why (a file in the old layout is told, not rejected).
RETIRED = {
    "MonthFundShifts": "fund shifts are made in Budgets",
    **{f"{p}Expenditures": "spend is read from the BCA inputs" for p in ("Month", "YTD")},
    **{
        f"{p}{kind}{unit}": "savings are estimated from the BCA inputs"
        for p in ("Month", "YTD")
        for kind in ("Gross", "Net")
        for unit in ("kWh", "kW", "Therm")
    },
}
SKIPPED = ("__MACOSX/", ".DS_Store", "Thumbs.db", "desktop.ini")


class UploadError(ValueError):
    """The upload is not a monthly report; the message says why, in words for the person who uploaded it."""


def _table_name(member: str) -> str | None:
    p = PurePosixPath(member)
    if len(p.parts) > 2 or p.suffix.lower() not in (".csv", ".txt"):
        return None
    return next((t for t in TABLES if p.stem.lower() == t.lower()), None)


def read_upload(data: bytes, filename: str) -> dict[str, bytes]:
    """{table: file bytes} from an uploaded .zip or .xlsx, every table present with its required columns."""
    if filename.lower().endswith(".xlsx"):
        tables = _from_workbook(data)
        _check_tables(tables, workbook=True)
        return tables
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise UploadError(
            "The file is not a valid .zip or .xlsx. Upload a .zip of Commitments.csv, Measure.csv and ProgramCost.csv, "
            "or an Excel workbook with those three sheets."
        ) from None
    tables, extra = {}, []
    for info in z.infolist():
        name = info.filename
        if info.is_dir() or any(s in name for s in SKIPPED) or PurePosixPath(name).name.startswith("._"):
            continue
        table = _table_name(name)
        if table is None:
            extra.append(name)
        elif table in tables:
            raise UploadError(f"The .zip holds two {table} files; keep one.")
        else:
            tables[table] = z.read(info)
    if extra:
        old = any(PurePosixPath(e).stem.lower() == "expenditures" for e in extra)
        raise UploadError(
            f"The .zip may hold only {', '.join(t + '.csv' for t in TABLES)}; it also has {', '.join(sorted(extra))}."
            + (" Spend is read from the BCA inputs (ProgramCost and the measures' incentives), so leave Expenditures out." if old else "")
        )
    _check_tables(tables)
    return tables


def _from_workbook(data: bytes) -> dict[str, bytes]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # openpyxl raises several kinds for a file that is not a workbook
        raise UploadError("The file is not a valid Excel workbook (.xlsx).") from None
    tables = {}
    for ws in wb.worksheets:
        table = next((t for t in TABLES if ws.title.strip().lower() == t.lower()), None)
        if table is None:
            continue
        out = io.StringIO()
        w = csv.writer(out, lineterminator="\n")
        for row in ws.iter_rows(values_only=True):
            if row is None or all(c is None or str(c).strip() == "" for c in row):
                continue
            w.writerow(["" if c is None else _cell(c) for c in row])
        tables[table] = out.getvalue().encode("utf-8")
    wb.close()
    return tables


def _cell(value) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _check_tables(tables: dict, workbook: bool = False) -> None:
    where = "sheet" if workbook else "file"
    missing = [t for t in TABLES if t not in tables]
    if missing:
        raise UploadError(
            f"There is no {' or '.join(missing)} {where}. A monthly report holds Commitments and the year-to-date BCA inputs Measure and "
            "ProgramCost (Measure may have no rows)."
        )
    for t in TABLES:
        header, _ = rows(tables[t])
        have = {h.lower() for h in header}
        lacking = [c for c in REQUIRED[t] if c.lower() not in have]
        if lacking:
            raise UploadError(f"{t} has no {', '.join(lacking)} column{'s' if len(lacking) > 1 else ''}.")


def job_zip(tables: dict[str, bytes]) -> bytes:
    """The BCA inputs as a job .zip, written the same way every time so the same tables give the same SHA-256."""
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for t in BCA_TABLES:
            info = zipfile.ZipInfo(f"{t}.csv", date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(info, tables[t])
    return out.getvalue()


def report_zip(tables: dict[str, bytes]) -> bytes:
    """All three tables as one .zip, as kept and published."""
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for t in TABLES:
            z.writestr(f"{t}.csv", tables[t])
    return out.getvalue()


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def text(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        return content.decode("cp1252")


def rows(content: bytes) -> tuple[list[str], list[dict[str, str]]]:
    """The header and the rows of one table, as text, keyed by the header (stripped). CEDARS' files are pipe-delimited; a tab or a
    comma in the header line is taken instead when it has no pipe."""
    body = text(content)
    first = body.split("\n", 1)[0]
    sep = "|" if "|" in first else ("\t" if "\t" in first else ",")
    reader = csv.reader(io.StringIO(body), delimiter=sep)
    try:
        header = [h.strip() for h in next(reader)]
    except StopIteration:
        return [], []
    out = []
    for r in reader:
        if not r or all(not c.strip() for c in r):
            continue
        out.append({h: (r[i].strip() if i < len(r) else "") for i, h in enumerate(header)})
    return header, out


def get(row: dict, field: str) -> str:
    """A row's value of a field, whatever the case of its column name."""
    if field in row:
        return row[field]
    low = field.lower()
    return next((v for k, v in row.items() if k.lower() == low), "")


def number(value: str, *, blank_ok: bool = True) -> Decimal | None:
    """A cell as a number: blank is 0 (or None when a value is required), anything else not a number is None."""
    v = (value or "").strip()
    if not v:
        return Decimal(0) if blank_ok else None
    try:
        d = Decimal(v)
    except InvalidOperation:
        return None
    return d if d.is_finite() else None


def year_month(value: str) -> tuple[int, int] | None:
    """CEDARS' YearMonth (202803) or 2028-03 as (year, month)."""
    v = (value or "").strip().replace("-", "")
    if len(v) != 6 or not v.isdigit() or not 1 <= int(v[4:]) <= 12:
        return None
    return int(v[:4]), int(v[4:])


def year_quarter(value: str) -> tuple[int, int] | None:
    """A ClaimYearQuarter (2028Q1) as (year, quarter)."""
    v = (value or "").strip().upper()
    if len(v) != 6 or not v[:4].isdigit() or v[4] != "Q" or v[5] not in "1234":
        return None
    return int(v[:4]), int(v[5])
