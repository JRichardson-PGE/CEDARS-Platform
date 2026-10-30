"""The PII vault's functions (decision G, BR-10.1, BR-10.4). A claims PII file (PII_Contact and PII_Site, as a .zip of CSVs or a workbook
with a sheet per table) is checked at once against CEDARS' PII data rules and the claim's own keys, and kept only when nothing is
rejected. Reading it needs claims.pii in the organisation (a data steward approves every such grant, and the CPUC can require a
federated session for it) or in the CPUC organisation. Every upload, rejection and read is audited with counts and keys, never values.

Nothing here is published, and no other part of the platform reads these tables: they join claims only through this service, by
ContactID and SiteID."""

import csv
import hashlib
import io
import re
import zipfile
from dataclasses import asdict, dataclass
from pathlib import PurePosixPath

from django.conf import settings
from django.db import transaction

from cedars.services import audit
from cedars.services.accounts import permissions as perms
from cedars.services.accounts import service as access
from cedars.services.accounts.service import AccessDenied

from . import spec
from .models import PIIContact, PIISite, PIIUpload
from .router import VAULT

DIGITS = re.compile(r"^\d+$")
SKIPPED = ("__MACOSX/", ".DS_Store", "Thumbs.db", "desktop.ini")


class PIIError(ValueError):
    """The file is not a PII file; the message says why without repeating anything in it."""


def enabled() -> bool:
    """Whether this environment has a vault (never the sandbox)."""
    return VAULT in settings.DATABASES


def can_read(user, organisation, assurance: str | None) -> bool:
    """claims.pii in the organisation, or in the CPUC organisation, in a session that meets the PII policy."""
    cpuc = access.cpuc()
    return access.has_permission(user, organisation, perms.PII, assurance) or (
        cpuc is not None and access.has_permission(user, cpuc, perms.PII, assurance)
    )


def can_upload(user, organisation, assurance: str | None) -> bool:
    return access.has_permission(user, organisation, perms.PII, assurance)


# --- reading a PII file


@dataclass
class Finding:
    table: str
    line: int
    key: str
    field: str
    rule: str
    message: str


def read(data: bytes, filename: str) -> dict[str, list[dict[str, str]]]:
    """{table: rows} with the specification's columns (others are dropped unread)."""
    raw = _from_workbook(data) if filename.lower().endswith(".xlsx") else _from_zip(data)
    if not raw:
        raise PIIError("The file holds neither PII_Contact nor PII_Site.")
    out = {}
    for name, rows in raw.items():
        if not rows:
            continue
        header = [h.strip().lower() for h in rows[0]]
        wanted = {c.lower(): c for c in spec.columns(name)}
        out[name] = [
            {wanted[h]: (r[i].strip() if i < len(r) else "") for i, h in enumerate(header) if h in wanted}
            for r in rows[1:]
            if any(c.strip() for c in r)
        ]
    return out


def _from_zip(data: bytes) -> dict[str, list[list[str]]]:
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise PIIError("The file is not a valid .zip or .xlsx of PII_Contact and PII_Site.") from None
    out = {}
    for info in z.infolist():
        name = info.filename
        if info.is_dir() or any(s in name for s in SKIPPED) or PurePosixPath(name).name.startswith("._"):
            continue
        p = PurePosixPath(name)
        table = spec.table_name(p.stem) if p.suffix.lower() in (".csv", ".txt") and len(p.parts) <= 2 else None
        if table is None:
            raise PIIError("The .zip may hold only PII_Contact and PII_Site (as CSV files); the claim records go in the claims upload.")
        if table in out:
            raise PIIError(f"The .zip holds two {table} files; keep one.")
        out[table] = _csv_rows(z.read(info))
    return out


def _from_workbook(data: bytes) -> dict[str, list[list[str]]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # openpyxl raises several kinds for a file that is not a workbook
        raise PIIError("The file is not a valid Excel workbook (.xlsx).") from None
    out = {}
    for ws in wb.worksheets:
        table = spec.table_name(ws.title)
        if table is None:
            continue
        out[table] = [["" if c is None else _cell(c) for c in row] for row in ws.iter_rows(values_only=True) if row is not None]
    wb.close()
    return out


def _cell(value) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _csv_rows(content: bytes) -> list[list[str]]:
    try:
        body = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        body = content.decode("cp1252")
    first = body.split("\n", 1)[0]
    sep = "|" if "|" in first else ("\t" if "\t" in first else ",")
    return [r for r in csv.reader(io.StringIO(body), delimiter=sep) if r]


# --- the checks


def check(tables: dict[str, list[dict]], code: str, known: dict[str, set[str]]) -> list[Finding]:
    """CEDARS' PII data rules and the links to the claim: every key starts with the PA code, is filled in, once; lengths, digits-only
    phone numbers, an @ in an email, an electric or a gas service account for each site; and each ContactID / SiteID is one the claim's
    ContactClaim / Site records name (`known`: {table: compared keys}). Findings carry the key and field, never a value."""
    out: list[Finding] = []
    for name, rows in tables.items():
        key = spec.KEYS[name]
        seen: dict[str, int] = {}
        for i, r in enumerate(rows, start=2):
            k = r.get(key, "")

            def add(fld, rule, message, k=k, i=i, name=name):
                out.append(Finding(name, i, k, fld, rule, message))

            if not k:
                add(key, "PII-REQ", f"{key} is blank")
                continue
            if not k.upper().startswith(code.upper()):
                add(key, "PII-ID", f"{key} does not start with {code}")
            if k.lower() in seen:
                add(key, "PII-KEY", f"{name} has this {key} on line {seen[k.lower()]} too")
            seen.setdefault(k.lower(), i)
            if k.strip().lower() not in known.get(name, set()):
                where = "ContactClaim" if name == "PII_Contact" else "Site"
                add(key, "PII-REF", f"No {where} record of this claim names this {key}")
            for f in spec.TABLES[name]:
                v = r.get(f.name, "")
                if not v:
                    continue
                if len(v) > f.length:
                    add(f.name, "PII-LEN", f"{f.name} is longer than {f.length} characters")
                if f.type == "phone" and not DIGITS.match(v):
                    add(f.name, "PII-NUM", f"{f.name} has characters other than digits")
                if f.type == "email" and "@" not in v:
                    add(f.name, "PII-EMAIL", f"{f.name} has no @")
            if name == "PII_Site" and not (r.get("ElectricServiceAccountID") or r.get("GasServiceAccountID")):
                add("ElectricServiceAccountID", "PII-ACCT", "Neither ElectricServiceAccountID nor GasServiceAccountID is filled in")
    return out


# --- storing and reading


def _require() -> None:
    if not enabled():
        raise AccessDenied("This environment has no PII vault, so it takes no personal information.")


def store(actor, organisation, version_id: int, data: bytes, filename: str, known: dict[str, set[str]], assurance: str | None):
    """Check a PII file and keep it as the version's PII when nothing fails; a rejected file keeps only its findings. The version's
    earlier accepted upload is superseded and its records deleted (an open version's PII is replaced, as its claims upload is)."""
    _require()
    if not can_upload(actor, organisation, assurance):
        raise AccessDenied(f"You need claims.pii in {organisation.code} to upload PII.")
    tables = read(data, filename)
    findings = check(tables, organisation.code, known)
    counts = {name: len(rows) for name, rows in tables.items()}
    sha = hashlib.sha256(data).hexdigest()
    with transaction.atomic(using=VAULT):
        up = PIIUpload.objects.create(
            organisation_id=organisation.pk,
            version_id=version_id,
            uploaded_by_id=getattr(actor, "pk", None),
            filename=filename[:255],
            sha256=sha,
            status=PIIUpload.Status.REJECTED if findings else PIIUpload.Status.ACCEPTED,
            counts=counts,
            findings=[asdict(f) for f in findings[:1000]],
        )
        if not findings:
            old = PIIUpload.objects.filter(version_id=version_id, status=PIIUpload.Status.ACCEPTED).exclude(pk=up.pk)
            PIIContact.objects.filter(upload__in=old).delete()
            PIISite.objects.filter(upload__in=old).delete()
            old.update(status=PIIUpload.Status.SUPERSEDED)
            PIIContact.objects.bulk_create(
                PIIContact(upload=up, contact_id=r["ContactID"], values={c: v for c, v in r.items() if c != "ContactID" and v})
                for r in tables.get("PII_Contact", [])
            )
            PIISite.objects.bulk_create(
                PIISite(upload=up, site_id=r["SiteID"], values={c: v for c, v in r.items() if c != "SiteID" and v})
                for r in tables.get("PII_Site", [])
            )
    audit.record(
        "vault.pii_accepted" if not findings else "vault.pii_rejected",
        object_type="vault.piiupload",
        object_id=str(up.pk),
        object_repr=f"PII for submission version {version_id}",
        after={"version": version_id, "counts": counts, "findings": len(findings), "sha256": sha},
        actor=actor,
        organisation=organisation,
    )
    return up


def latest(version_ids: list[int]) -> PIIUpload | None:
    """The newest upload of the first of these versions that has one (a reopened version reads its predecessor's PII until its own)."""
    for vid in version_ids:
        up = PIIUpload.objects.filter(version_id=vid).exclude(status=PIIUpload.Status.SUPERSEDED).first()
        if up is not None:
            return up
    return None


def accepted(version_ids: list[int]) -> PIIUpload | None:
    for vid in version_ids:
        up = PIIUpload.objects.filter(version_id=vid, status=PIIUpload.Status.ACCEPTED).first()
        if up is not None:
            return up
    return None


def export(actor, organisation, up: PIIUpload, assurance: str | None, reason: str = "") -> bytes:
    """An accepted upload's records as a .zip of PII_Contact.csv and PII_Site.csv, for a reader with the PII grant; audited."""
    _require()
    if not can_read(actor, organisation, assurance):
        raise AccessDenied("Reading PII needs the PII grant, approved by a data steward.")
    if up.organisation_id != organisation.pk or up.status != PIIUpload.Status.ACCEPTED:
        raise AccessDenied("That is not this organisation's PII.")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        tables = (("PII_Contact", up.contacts.order_by("contact_id"), "contact_id"), ("PII_Site", up.sites.order_by("site_id"), "site_id"))
        for name, rows, key in tables:
            cols = spec.columns(name)
            buf = io.StringIO()
            w = csv.writer(buf, lineterminator="\n")
            w.writerow(cols)
            for r in rows:
                w.writerow([getattr(r, key)] + [r.values.get(c, "") for c in cols[1:]])
            z.writestr(f"{name}.csv", buf.getvalue())
    audit.record(
        "vault.pii_read",
        object_type="vault.piiupload",
        object_id=str(up.pk),
        object_repr=f"PII for submission version {up.version_id}",
        reason=reason,
        after={"version": up.version_id, "counts": up.counts},
        actor=actor,
        organisation=organisation,
    )
    return out.getvalue()
