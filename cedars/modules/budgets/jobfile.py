"""A budget filing's upload: a .zip holding Measure.csv and ProgramCost.csv (as CET takes a job; comma-, pipe- or tab-delimited), or
an Excel workbook with Measure and ProgramCost sheets (BR-02.2). A workbook becomes a job .zip of two comma-delimited CSVs, so the
version's job is always a .zip the engine reads; a .zip is kept exactly as uploaded."""

import csv
import io
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import PurePosixPath

TABLES = ("Measure", "ProgramCost")
REQUIRED = {
    "Measure": ["CEInputID", "PrgID", "ClaimYearQuarter", "NumUnits"],
    "ProgramCost": ["PrgID", "PrgYear", "PA"],
}
SKIPPED = ("__MACOSX/", ".DS_Store", "Thumbs.db", "desktop.ini")


class JobError(ValueError):
    """The upload is not a budget filing job; the message says why, in words for the person who uploaded it."""


def _table_name(member: str) -> str | None:
    p = PurePosixPath(member)
    if len(p.parts) > 2 or p.suffix.lower() not in (".csv", ".txt"):
        return None
    return next((t for t in TABLES if p.stem.lower() == t.lower()), None)


def read_upload(data: bytes, filename: str) -> tuple[dict[str, bytes], bytes]:
    """({table: file bytes}, the job .zip) from an uploaded .zip or .xlsx."""
    if filename.lower().endswith(".xlsx"):
        tables = _from_workbook(data)
        return tables, job_zip(tables)
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise JobError(
            "The file is not a valid .zip or .xlsx. Upload a .zip of Measure.csv and ProgramCost.csv, or an Excel workbook."
        ) from None
    tables, extra = {}, []
    for info in z.infolist():
        name = info.filename
        if info.is_dir() or any(s in name for s in SKIPPED) or PurePosixPath(name).name.startswith("._"):
            continue
        table = _table_name(name)
        if table is None:
            extra.append(name)
        elif table in tables:
            raise JobError(f"The .zip holds two {table} files; keep one.")
        else:
            tables[table] = z.read(info)
    if extra:
        raise JobError(f"The .zip may hold only Measure.csv and ProgramCost.csv; it also has {', '.join(sorted(extra))}.")
    _check_tables(tables)
    return tables, data


def _from_workbook(data: bytes) -> dict[str, bytes]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # openpyxl raises several kinds for a file that is not a workbook
        raise JobError("The file is not a valid Excel workbook (.xlsx).") from None
    tables = {}
    for ws in wb.worksheets:
        table = next((t for t in TABLES if ws.title.strip().lower() == t.lower()), None)
        if table is None:
            continue
        out = io.StringIO()
        w = csv.writer(out, lineterminator="\n")
        for row in ws.iter_rows(values_only=True):
            if row is None or all(c is None or str(c).strip() == "" for c in row):
                continue
            w.writerow(["" if c is None else _cell(c) for c in row])
        tables[table] = out.getvalue().encode("utf-8")
    wb.close()
    _check_tables(tables, workbook=True)
    return tables


def _cell(value) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _check_tables(tables: dict, workbook: bool = False) -> None:
    where = "sheet" if workbook else "file"
    for t in TABLES:
        if t not in tables:
            raise JobError(f"There is no {t} {where}. A budget filing needs both Measure and ProgramCost.")
    for t in TABLES:
        header, _ = rows(tables[t])
        have = {h.lower() for h in header}
        missing = [c for c in REQUIRED[t] if c.lower() not in have]
        if missing:
            raise JobError(f"{t} has no {', '.join(missing)} column{'s' if len(missing) > 1 else ''}.")


def job_zip(tables: dict[str, bytes]) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for t in TABLES:
            z.writestr(f"{t}.csv", tables[t])
    return out.getvalue()


def text(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        return content.decode("cp1252")


def rows(content: bytes) -> tuple[list[str], list[dict[str, str]]]:
    """The header and the rows of one table, as text, keyed by the header (stripped). CET's files are pipe-delimited; a tab or a comma
    in the header line is taken instead when it has no pipe."""
    body = text(content)
    first = body.split("\n", 1)[0]
    sep = "|" if "|" in first else ("\t" if "\t" in first else ",")
    reader = csv.reader(io.StringIO(body), delimiter=sep)
    try:
        header = [h.strip() for h in next(reader)]
    except StopIteration:
        return [], []
    out = []
    for r in reader:
        if not r or all(not c.strip() for c in r):
            continue
        out.append({h: (r[i].strip() if i < len(r) else "") for i, h in enumerate(header)})
    return header, out


def get(row: dict, field: str) -> str:
    """A row's value of a field, whatever the case of its column name."""
    if field in row:
        return row[field]
    low = field.lower()
    return next((v for k, v in row.items() if k.lower() == low), "")


def number(value: str) -> Decimal | None:
    """A cell as a number: blank is 0, anything else not a number is None."""
    v = (value or "").strip()
    if not v:
        return Decimal(0)
    try:
        d = Decimal(v)
    except InvalidOperation:
        return None
    return d if d.is_finite() else None


def year_of(value: str) -> int | None:
    v = (value or "").strip()
    return int(v[:4]) if len(v) >= 4 and v[:4].isdigit() else None
