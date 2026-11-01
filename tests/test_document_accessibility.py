"""Accessible documents (BRD v0.5 BR-10.3): every generated workbook passes the document accessibility check, and uploaded PDFs and
workbooks are checked, the uploader warned, the files listed for review, and refused only when the CPUC's settings say so."""

import io

import pyarrow as pa
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from pypdf import PdfWriter
from pypdf.generic import ArrayObject, BooleanObject, DecodedStreamObject, DictionaryObject, NameObject, TextStringObject

from cedars.core import workbooks
from cedars.modules.claims import spec
from cedars.modules.data import downloads, records
from cedars.modules.documents import service as documents
from cedars.modules.documents.models import DocumentVersion
from cedars.services.accessibility import checks
from cedars.services.accessibility import service as accessibility
from cedars.services.accessibility.models import FileCheck
from cedars.services.accounts.service import AccessDenied
from cedars.services.audit.models import AuditEvent
from cedars.services.submissions import comments as filing_comments

from .conftest import page_only
from .test_documents import details, make_people
from .test_submission_comments import _priced, analyst, approver, deadline  # noqa: F401  (fixtures)

pytestmark = pytest.mark.django_db


def pdf(*, text=True, title="2025 Annual Report", lang="en-US", tagged=True, password="") -> bytes:
    """A one-page PDF with what an accessible PDF has (tagged text, a title, a language), each of which can be left out."""
    w = PdfWriter()
    page = w.add_blank_page(612, 792)
    if text:
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): w._add_object(font)})})
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 72 720 Td (Energy efficiency annual report) Tj ET")
        page[NameObject("/Contents")] = w._add_object(content)
    if title:
        w.add_metadata({"/Title": title})
    root = w._root_object
    if lang:
        root[NameObject("/Lang")] = TextStringObject(lang)
    if tagged:
        root[NameObject("/MarkInfo")] = DictionaryObject({NameObject("/Marked"): BooleanObject(True)})
        tree = DictionaryObject({NameObject("/Type"): NameObject("/StructTreeRoot"), NameObject("/K"): ArrayObject()})
        root[NameObject("/StructTreeRoot")] = w._add_object(tree)
    if password:
        w.encrypt(password)
    out = io.BytesIO()
    w.write(out)
    return out.getvalue()


def plain_workbook(merge=False, chart=False, name="Sheet") -> bytes:
    """A workbook as people make one: a bold-looking first row that nothing marks as a header."""
    wb = Workbook()
    ws = wb.active
    ws.title = name
    ws.append(["Program", "Budget"])
    ws.append(["PGE_Res_001", 1000])
    if merge:
        ws.merge_cells("A4:B4")
    if chart:
        from openpyxl.chart import BarChart, Reference

        c = BarChart()
        c.add_data(Reference(ws, min_col=2, min_row=1, max_row=2), titles_from_data=True)
        ws.add_chart(c, "D2")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def codes(data: bytes, name: str) -> set[str]:
    return checks.check_bytes(data, name).codes()


# --- the check


def test_an_accessible_pdf_has_no_findings_and_each_missing_part_is_found():
    assert checks.check_bytes(pdf(), "report.pdf").clear
    assert codes(pdf(tagged=False), "r.pdf") == {"pdf_untagged"}
    assert codes(pdf(text=False), "r.pdf") == {"pdf_image_only"}
    assert codes(pdf(title=""), "r.pdf") == {"pdf_no_title"}
    assert codes(pdf(lang=""), "r.pdf") == {"pdf_no_language"}
    assert codes(pdf(text=False, title="", lang="", tagged=False), "scan.PDF") == set(c for c in checks.CODES if c.startswith("pdf_"))


def test_a_pdf_that_cannot_be_read_is_noted_not_found_wanting():
    locked = checks.check_bytes(pdf(password="secret"), "locked.pdf")
    assert locked.kind == "pdf" and not locked.findings and "password" in locked.note
    broken = checks.check_bytes(b"%PDF-1.4 not really", "broken.pdf")
    assert not broken.findings and broken.note and not broken.checked


def test_a_workbook_needs_marked_headers_no_merges_alt_text_and_sheet_names():
    assert codes(plain_workbook(), "w.xlsx") == {"workbook_no_headers", "workbook_default_sheet_name"}
    assert codes(plain_workbook(name="Budget", merge=True), "w.xlsx") == {"workbook_no_headers", "workbook_merged_cells"}
    found = checks.check_bytes(plain_workbook(name="Budget", chart=True), "w.xlsm").findings
    assert [f.code for f in found] == ["workbook_no_headers", "workbook_no_alt_text"] and found[1].where == "Budget"


def test_print_titles_or_a_title_name_mark_a_header_row_as_a_table_does():
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Program", "Budget"])
    ws.append(["PGE_Res_001", 1000])
    ws.print_title_rows = "1:1"
    out = io.BytesIO()
    wb.save(out)
    assert codes(out.getvalue(), "w.xlsx") == set()


def test_other_types_are_not_checked():
    r = checks.check_bytes(b"hello", "notes.txt")
    assert not r.kind and not r.findings and not checks.checkable("deck.pptx") and checks.checkable("book.XLSX")


# --- generated workbooks


def assert_accessible(data: bytes) -> None:
    result = checks.check_bytes(data, "generated.xlsx")
    assert result.checked and result.clear, [f.describe() for f in result.findings]
    wb = load_workbook(io.BytesIO(data))
    assert wb.properties.title and wb.properties.language == "en-US"
    for ws in wb.worksheets:
        assert ws.tables and not ws.merged_cells.ranges, ws.title
        (table,) = ws.tables.values()
        assert table.headerRowCount == 1 and [c.name for c in table.tableColumns] == [c.value for c in ws[1]]
        assert ws.freeze_panes == "A2" and ws.print_title_rows == "$1:$1"


def test_every_generated_workbook_passes_the_check():
    assert_accessible(spec.template_xlsx())
    assert_accessible(downloads._xlsx(["PA", "TRC ratio", "Records"], [["PGE", 1.2, 10], ["Total", 1.2, 10]], {"dataset": "claims"}))
    assert_accessible(downloads._xlsx(["PA", "Records"], [], {"dataset": "claims"}))  # no rows: the table keeps an empty one
    batches = [pa.record_batch({"CEInputID": ["S1", "S2"], "TotalSystemBenefit": [11, 22]})]
    out = io.BytesIO()
    assert records._xlsx(batches, batches[0].schema, {"dataset": "claims", "records": 2}, out) == 2
    assert_accessible(out.getvalue())


def test_the_workbook_writer_makes_headers_a_table_can_take():
    book = workbooks.Book("Test")
    book.sheet("Data", ["PrgID", "prgid", None, ""], [[1, 2, 3, 4]])
    book.sheet("Data copy", ["PrgID"])
    data = book.save()
    assert_accessible(data)
    wb = load_workbook(io.BytesIO(data))
    assert [c.value for c in wb["Data"][1]] == ["PrgID", "prgid (2)", "Column 3", "Column 4"]
    assert sorted(t for ws in wb.worksheets for t in ws.tables) == ["DataCopyTable", "DataTable"]


# --- uploads


@pytest.fixture
def people(cpuc_admin, cpuc, pge, pge_admin, make_user):
    return make_people(cpuc_admin, cpuc, pge, pge_admin, make_user)


def file_document(people, pge, *files):
    uploads = [documents.Upload(n, SimpleUploadedFile(n, d)) for n, d in files]
    return documents.create_document(people["pge_writer"], pge, "pa", details("pa", "Joint Cooperation Memo"), uploads)


def test_an_upload_is_checked_kept_and_audited(people, pge):
    doc = file_document(
        people, pge, ("memo.pdf", pdf(tagged=False, lang="")), ("tables.xlsx", plain_workbook(name="Budget")), ("a.txt", b"x")
    )
    main, book, _ = doc.versions.get().files.all()
    checked = accessibility.for_refs("document", [main.pk, book.pk])
    assert set(checked) == {str(main.pk), str(book.pk)}  # the text file is not a checked type
    fc = checked[str(main.pk)]
    assert {f.code for f in fc.finding_list} == {"pdf_untagged", "pdf_no_language"}
    assert (fc.organisation, fc.uploaded_by, fc.url, fc.sha256) == (pge, people["pge_writer"], doc.get_absolute_url(), main.sha256)
    assert AuditEvent.objects.filter(action="accessibility.checked", object_id=str(fc.pk)).exists()
    warning = accessibility.upload_warning("document", [main.pk, book.pk])
    assert (
        "memo.pdf: PDF is not tagged; PDF has no language" in warning and "tables.xlsx: Sheet has no marked header row (Budget)" in warning
    )


def test_a_clean_upload_gets_no_warning(people, pge):
    doc = file_document(people, pge, ("report.pdf", pdf()))
    (f,) = doc.versions.get().files.all()
    assert accessibility.for_refs("document", [f.pk])[str(f.pk)].findings == []
    assert accessibility.upload_warning("document", [f.pk]) == ""


def test_nothing_is_refused_until_the_cpuc_turns_blocking_on(people, pge, cpuc_admin):
    file_document(people, pge, ("scan.pdf", pdf(text=False)))
    with pytest.raises(AccessDenied):
        accessibility.set_settings(people["pge_writer"], block=True, blocking=["pdf_image_only"], reason="x")
    with pytest.raises(AccessDenied):
        accessibility.set_settings(cpuc_admin, block=True, blocking=["pdf_image_only"], reason="")
    with pytest.raises(AccessDenied):
        accessibility.set_settings(cpuc_admin, block=True, blocking=[], reason="Onboarding")
    accessibility.set_settings(cpuc_admin, block=True, blocking=["pdf_image_only", "nonsense"], reason="PA onboarding")
    assert accessibility.settings_row().blocking == ["pdf_image_only"]
    assert AuditEvent.objects.filter(action="accessibility.settings_changed", reason="PA onboarding").exists()
    with pytest.raises(documents.DocumentError, match="scan.pdf is not accessible enough to be filed: PDF has no text"):
        file_document(people, pge, ("scan.pdf", pdf(text=False)))
    file_document(people, pge, ("untagged.pdf", pdf(tagged=False)))  # a finding the settings do not block


def test_a_legacy_import_is_listed_but_never_refused(people, pge, cpuc_admin):
    accessibility.set_settings(cpuc_admin, block=True, blocking=["pdf_image_only"], reason="PA onboarding")
    doc = file_document(people, pge, ("first.pdf", pdf()))
    uploads = [documents.Upload("scan.pdf", SimpleUploadedFile("scan.pdf", pdf(text=False)))]
    version = DocumentVersion.objects.create(document=doc, number=2, reason="Imported from the PDA", organisation=pge)
    files = documents._store_files(version, uploads, documents._check_uploads(uploads), legacy_names=["scan.pdf"])
    (fc,) = accessibility.for_refs("document", [files[0].pk]).values()
    assert fc.imported and [f.code for f in fc.finding_list] == ["pdf_image_only"]


def test_a_file_on_a_filing_comment_is_checked(analyst, approver, deadline, pge, cpuc_admin):  # noqa: F811
    v = _priced(analyst, deadline, pge)
    c = filing_comments.post(approver, v, "See the table.", [("table.xlsx", io.BytesIO(plain_workbook(name="NTG")))])
    (a,) = c.attachments.all()
    fc = FileCheck.objects.get(source="filing_comment", ref=str(a.pk))
    assert [f.code for f in fc.finding_list] == ["workbook_no_headers"] and fc.uploaded_by == approver and fc.url
    accessibility.set_settings(cpuc_admin, block=True, blocking=["workbook_no_headers"], reason="PA onboarding")
    with pytest.raises(filing_comments.CommentRefused, match="table.xlsx is not accessible enough"):
        filing_comments.post(approver, v, "Again.", [("table.xlsx", io.BytesIO(plain_workbook(name="NTG")))])


# --- pages


def test_the_document_page_says_a_file_may_not_be_accessible(client, signed_in, people, pge):
    doc = file_document(people, pge, ("memo.pdf", pdf(lang="")), ("ok.pdf", pdf()))
    signed_in(people["pge_writer"])
    page = page_only(client.get(doc.get_absolute_url()).content.decode())
    assert page.count("May not be accessible") == 1 and "PDF has no language" in page and "ask for an accessible copy" in page


def test_the_admin_list_review_and_settings(client, signed_in, people, pge, cpuc_admin):
    file_document(people, pge, ("memo.pdf", pdf(lang="")))
    fc = FileCheck.objects.get()
    signed_in(people["pge_writer"])
    assert client.get(reverse("accessibility:files")).status_code == 404
    assert client.get(reverse("accessibility:settings")).status_code == 404
    signed_in(cpuc_admin)
    page = client.get(reverse("accessibility:files")).content.decode()
    assert "memo.pdf" in page and "Needs review" in page and "Warning only" in page
    assert "memo.pdf" not in page_only(client.get(reverse("accessibility:files") + "?state=clear").content.decode())
    assert "memo.pdf" in client.get(reverse("accessibility:files") + "?code=pdf_no_language&kind=pdf").content.decode()
    assert "memo.pdf" not in page_only(client.get(reverse("accessibility:files") + "?code=pdf_untagged").content.decode())
    detail = client.get(fc.get_absolute_url()).content.decode()
    assert "PDF has no language" in detail and "3.1.1" in detail
    client.post(reverse("accessibility:review", args=[fc.pk]), {"note": ""})
    fc.refresh_from_db()
    assert not fc.reviewed
    client.post(reverse("accessibility:review", args=[fc.pk]), {"note": "Accessible copy filed as version 2"})
    fc.refresh_from_db()
    assert fc.reviewed and fc.reviewed_by == cpuc_admin
    assert "Needs review" not in page_only(client.get(reverse("accessibility:files")).content.decode())
    client.post(reverse("accessibility:reopen", args=[fc.pk]), {"reason": "Version 2 is untagged too"})
    fc.refresh_from_db()
    assert not fc.reviewed
    assert [e.action for e in AuditEvent.objects.filter(object_id=str(fc.pk), action__startswith="accessibility.").order_by("id")] == [
        "accessibility.checked",
        "accessibility.reviewed",
        "accessibility.reopened",
    ]
    client.post(reverse("accessibility:settings"), {"block": "on", "blocking": ["pdf_no_title"], "reason": "Onboarding"})
    cfg = accessibility.settings_row()
    assert cfg.block and cfg.blocking == ["pdf_no_title"]
    assert "Blocking on" in client.get(reverse("accessibility:files")).content.decode()
