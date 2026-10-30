"""Budget filings (BR-04). A filing is a deadline of the "budgets" module on a filing cycle (business plan application, TUAL, MCAL;
the submissions service's calendar) with this module's FilingSetup; each PA's filing is a submission with versions.

    upload (.zip or .xlsx) -> lines and checks -> validated -> priced (Calculations, through the submissions service) -> confirm
    -> published: inputs, run settings and results readable by every signed-in user (BR-04.4)

Every filed version is kept and comparable (BR-04.2); results are read by program year and totalled over the years chosen (BR-04.1);
programs flagged exclude-from-budget or exclude-from-C/E keep their values and are filtered, never zeroed (BR-04.3); each version
records the allocation version it was priced with, and a later one never changes it (BR-04.8)."""

import hashlib
import io
import json
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime, time
from decimal import Decimal
from pathlib import Path

from django.db import transaction
from django.utils import timezone

from cedars.core import storage
from cedars.services import audit
from cedars.services.accounts import service as access
from cedars.services.accounts.service import AccessDenied
from cedars.services.registry import service as registry
from cedars.services.registry.models import AllocationVersion, Program
from cedars.services.rules import service as rules
from cedars.services.submissions import service as submissions
from cedars.services.submissions.models import Deadline, Submission, SubmissionVersion

from . import jobfile
from .models import PROGRAM_COST_FIELDS, BudgetLine, BudgetResult, FilingSetup

MODULE = "budgets"
RULE_SET = "budget_filing"
INCENTIVE_FIELDS = ["UnitIncentiveToOthers", "UnitDirectInstallLab", "UnitDirectInstallMat", "UnitEndUserRebate"]
RESULT_KEYS = {"_pa", "_prg", "_prgname", "_sector", "_segment", "_year", "_n"}


# --- filings (the CPUC sets them up)


def filings(cycle=None):
    qs = Deadline.objects.filter(module=MODULE, cycle__isnull=False).select_related("cycle", "budget_setup__allocation_version")
    return qs.filter(cycle=cycle) if cycle is not None else qs


def label(deadline: Deadline) -> str:
    return f"{deadline.cycle} {deadline.kind}"


@transaction.atomic
def add_filing(
    actor, *, first_year: int, kind: str, due_at, acc_version: str, allocation_version=None, opens_at=None, notes=""
) -> Deadline:
    """A budget filing for a four-year cycle, due on a date, priced on an ACC (and an allocation version, or the most recent)."""
    if actor is not None and not access.is_system_admin(actor):
        raise AccessDenied("Only a CPUC system administrator sets up budget filings.")
    _check_allocation(allocation_version)
    cycle = registry.ensure_cycle(first_year)
    if Deadline.objects.filter(module=MODULE, kind=kind, cycle=cycle).exists():
        raise AccessDenied(f"The {cycle} cycle already has a {kind} filing.")
    deadline = submissions.add_deadline(actor, module=MODULE, kind=kind, cycle=cycle, due_at=due_at, opens_at=opens_at, notes=notes)
    setup = FilingSetup.objects.create(deadline=deadline, acc_version=acc_version, allocation_version=allocation_version)
    audit.record("budget_filing.set_up", obj=setup, after=_setup_fields(setup), actor=actor)
    return deadline


def _setup_fields(setup: FilingSetup) -> dict:
    return {"acc": setup.acc_version, "allocation_version": str(setup.allocation_version or "most recent")}


def _check_allocation(version: AllocationVersion | None) -> None:
    if version is not None and version.status != AllocationVersion.Status.CONFIRMED:
        raise AccessDenied(f"{version} is a draft; a filing is priced only with a confirmed allocation version.")


@transaction.atomic
def change_setup(actor, setup: FilingSetup, *, acc_version: str, allocation_version=None, reason: str) -> FilingSetup:
    """Change how a filing is priced. Versions already priced keep what they were priced with; the next pricing uses the new setting."""
    if actor is not None and not access.is_system_admin(actor):
        raise AccessDenied("Only a CPUC system administrator changes how a budget filing is priced.")
    if not reason.strip():
        raise AccessDenied("Give the reason for the change.")
    _check_allocation(allocation_version)
    before = _setup_fields(setup)
    setup.acc_version, setup.allocation_version = acc_version, allocation_version
    setup.save()
    audit.record("budget_filing.setup_changed", obj=setup, before=before, after=_setup_fields(setup), reason=reason, actor=actor)
    return setup


def due_at(day) -> datetime:
    """A due date as the end of that day."""
    return timezone.make_aware(datetime.combine(day, time(23, 59, 59)))


# --- who sees what


def _org_reader(user, organisation) -> bool:
    return any(access.has_permission(user, organisation, f"{MODULE}.{a}") for a in ("read", "submit", "confirm"))


def _cpuc_reader(user) -> bool:
    org = access.cpuc()
    return access.is_system_admin(user) or (
        org is not None and any(access.has_permission(user, org, f"{MODULE}.{a}") for a in ("read", "approve"))
    )


def can_read_submission(user, submission: Submission) -> bool:
    return _org_reader(user, submission.organisation) or _cpuc_reader(user)


def can_read(user, v: SubmissionVersion) -> bool:
    """A published version is public to every signed-in user (BR-04.4); the rest only to the organisation and CPUC staff."""
    if not user.is_authenticated:
        return False
    return v.status == SubmissionVersion.Status.PUBLISHED or can_read_submission(user, v.submission)


def readable_versions(user):
    """Every version the user may read, newest filing first."""
    qs = SubmissionVersion.objects.filter(submission__deadline__module=MODULE).select_related(
        "submission__deadline__cycle", "submission__organisation"
    )
    if not _cpuc_reader(user):
        orgs = [o for o in access.organisations_of(user) if _org_reader(user, o)]
        qs = qs.filter(status=SubmissionVersion.Status.PUBLISHED) | qs.filter(submission__organisation__in=orgs)
    return qs.order_by("-submission__deadline__due_at", "submission__organisation__code", "-number")


def standing(organisation, cycle) -> SubmissionVersion | None:
    """The budget that stands for a PA's cycle: the published version of its latest filing (BR-04.2: the most recent by default)."""
    return (
        SubmissionVersion.objects.filter(
            submission__organisation=organisation,
            submission__deadline__module=MODULE,
            submission__deadline__cycle=cycle,
            status=SubmissionVersion.Status.PUBLISHED,
        )
        .select_related("submission__deadline__cycle")
        .order_by("-submission__deadline__due_at", "-number")
        .first()
    )


# --- the upload and its checks


@dataclass
class Problem:
    """A module check's finding, in the rule library's format (rules.Finding)."""

    row: int
    record_id: str
    table: str
    field: str
    rule_id: str
    severity: str
    message: str
    fix: str
    value: str = ""


def _prefix(v: SubmissionVersion) -> str:
    return f"budgets/{v.submission_id}/v{v.number}/"


def derive(tables: dict[str, bytes], organisation, cycle) -> tuple[list[dict], list[Problem], int]:
    """The version's budget lines {prg_id, year, pa, program_costs, on_bill_financing, rebates_and_incentives, measures} from its
    tables, the module's own checks, and the number of records. Module checks: each ProgramCost year and measure claim year in the
    cycle, every PrgID in the program registry and the organisation's own (or statewide), the cost and unit fields numbers, and a
    warning for measures whose program has no ProgramCost row that year (priced with no program costs)."""
    problems: list[Problem] = []
    _, costs = jobfile.rows(tables["ProgramCost"])
    _, measures = jobfile.rows(tables["Measure"])
    programs = {p.prg_id: p for p in Program.objects.filter(prg_id__in={jobfile.get(r, "PrgID") for r in costs + measures})}
    lines: dict[tuple[str, int], dict] = {}

    def program_ok(table, i, rid, prg) -> bool:
        p = programs.get(prg)
        if p is None:
            problems.append(
                Problem(
                    i,
                    rid,
                    table,
                    "PrgID",
                    "BUD-PRG",
                    "reject",
                    f"PrgID {prg or '(blank)'} is not in the program registry",
                    "Use the PrgID of a program in the Programs catalog, or ask for the program to be added first.",
                    prg,
                )
            )
            return False
        if not p.statewide and p.administrator_id not in (None, organisation.pk):
            problems.append(
                Problem(
                    i,
                    rid,
                    table,
                    "PrgID",
                    "BUD-ORG",
                    "reject",
                    f"{prg} is another program administrator's program",
                    f"File only {organisation.code}'s programs and the statewide programs it leads.",
                    prg,
                )
            )
            return False
        return True

    def year_ok(table, i, rid, fieldname, year, raw) -> bool:
        if year is None or not cycle.covers(year):
            problems.append(
                Problem(
                    i,
                    rid,
                    table,
                    fieldname,
                    "BUD-YEAR",
                    "reject",
                    f"{fieldname} {raw or '(blank)'} is not a year of the {cycle} cycle",
                    f"Use a year from {cycle.first_year} to {cycle.last_year}.",
                    raw,
                )
            )
            return False
        return True

    for i, r in enumerate(costs, start=2):
        prg, raw = jobfile.get(r, "PrgID"), jobfile.get(r, "PrgYear")
        year = jobfile.year_of(raw) if len(raw.strip()) == 4 else None
        ok = program_ok("ProgramCost", i, prg, prg) & year_ok("ProgramCost", i, prg, "PrgYear", year, raw)
        values = {}
        for f, _ in PROGRAM_COST_FIELDS + [("OnBillFinancing", "")]:
            n = jobfile.number(jobfile.get(r, f))
            if n is None:
                problems.append(
                    Problem(
                        i,
                        prg,
                        "ProgramCost",
                        f,
                        "BUD-NUM",
                        "reject",
                        f"{f} should be a number",
                        f"Enter a plain number in {f}.",
                        jobfile.get(r, f),
                    )
                )
                ok = False
            values[f] = n or Decimal(0)
        if not ok:
            continue
        if (prg, year) in lines:
            problems.append(
                Problem(
                    i,
                    prg,
                    "ProgramCost",
                    "PrgYear",
                    "BUD-DUP",
                    "reject",
                    f"{prg} has two ProgramCost rows for {year}",
                    "Keep one row per program and year.",
                    str(year),
                )
            )
            continue
        lines[(prg, year)] = {
            "prg_id": prg,
            "year": year,
            "pa": jobfile.get(r, "PA"),
            "program_costs": {f: str(values[f]) for f, _ in PROGRAM_COST_FIELDS},
            "on_bill_financing": values["OnBillFinancing"],
            "rebates_and_incentives": Decimal(0),
            "measures": 0,
        }
    without_costs: dict[tuple[str, int], int] = defaultdict(int)
    for i, r in enumerate(measures, start=2):
        rid, prg, raw = jobfile.get(r, "CEInputID"), jobfile.get(r, "PrgID"), jobfile.get(r, "ClaimYearQuarter")
        year = jobfile.year_of(raw)
        if not (program_ok("Measure", i, rid, prg) & year_ok("Measure", i, rid, "ClaimYearQuarter", year, raw)):
            continue
        units = jobfile.number(jobfile.get(r, "NumUnits"))
        unit = [jobfile.number(jobfile.get(r, f)) for f in INCENTIVE_FIELDS]
        bad = (["NumUnits"] if units is None else []) + [f for f, n in zip(INCENTIVE_FIELDS, unit, strict=True) if n is None]
        for f in bad:
            problems.append(
                Problem(
                    i, rid, "Measure", f, "BUD-NUM", "reject", f"{f} should be a number", f"Enter a plain number in {f}.", jobfile.get(r, f)
                )
            )
        if bad:
            continue
        line = lines.get((prg, year))
        if line is None:
            without_costs[(prg, year)] += 1
            line = lines[(prg, year)] = {
                "prg_id": prg,
                "year": year,
                "pa": "",
                "program_costs": {},
                "on_bill_financing": Decimal(0),
                "rebates_and_incentives": Decimal(0),
                "measures": 0,
            }
        line["rebates_and_incentives"] += units * sum(unit)
        line["measures"] += 1
    for (prg, year), n in sorted(without_costs.items()):
        problems.append(
            Problem(
                0,
                prg,
                "Measure",
                "PrgID",
                "BUD-NOCOST",
                "warning",
                f"{n} measure{'s' if n > 1 else ''} of {prg} in {year} {'have' if n > 1 else 'has'} no ProgramCost row for that year: "
                "priced with no program costs",
                f"Add {prg}'s {year} ProgramCost row if it has program costs.",
                str(year),
            )
        )
    return list(lines.values()), problems, len(costs) + len(measures)


def _rule_findings(tables: dict[str, bytes], cycle) -> tuple[list, list[str], str]:
    """The budget filing rule set in force (CEDARS' own rules, run by the engine), year by year: CEDARS' rules take a single filing
    year, and a cycle filing spans four. Findings keep the file's own line numbers."""
    rs = rules.in_force(RULE_SET)
    if rs is None:
        return [], ["No budget filing rule set is published yet, so only the module's own checks ran."], ""
    try:
        frames = {t: rules.read_table(tables[t]) for t in jobfile.TABLES}
    except rules.EngineMissing as e:
        raise AccessDenied(f"{e} The filing cannot be validated on this server.") from None
    out, notes = [], set()
    for table, year_col, id_col in (("ProgramCost", "PrgYear", "PrgID"), ("Measure", "ClaimYearQuarter", "CEInputID")):
        df = frames[table]
        have = {c.lower() for c in df.columns}
        spec = sorted(set(rs.rules.filter(table=table, kind="spec").values_list("field", flat=True)))
        missing = [c for c in spec if c.lower() not in have and _recet_name(c).lower() not in have]
        if missing:  # CEDARS rejects a file without its spec's columns; the rules are run once they are there
            out += [
                Problem(1, "", table, c, "BUD-COL", "reject", f"{table} has no {c} column", f"Add the {c} column (it may be blank).")
                for c in missing
            ]
            notes.add(f"The {table} rules were not run: the file lacks {len(missing)} of the filing specification's columns.")
            continue
        if year_col not in df.columns:
            continue
        years = df[year_col].astype(str).str[:4]
        for y in sorted(set(years)):
            if not (y.isdigit() and cycle.covers(int(y))):
                continue  # the module's own check reports it
            part = df[years == y]
            try:
                found, n = rules.check(rs, table, part.reset_index(drop=True), filing_year=int(y), id_column=id_col)
            except (KeyError, ValueError) as e:  # the engine could not apply a rule: the version cannot be called valid
                out.append(
                    Problem(1, "", table, "", "BUD-RULES", "reject", f"The {y} {table} rows could not be checked: {e}",
                            "Tell the CPUC's CEDARS support; the file may be fine.")
                )  # fmt: skip
                continue
            lines = list(part.index)
            for f in found:
                f.row = int(lines[f.row - 2]) + 2
            out += found
            notes.update(n)
    return out, sorted(notes), f"{rs.label} ({rs.effective_from})"


def _recet_name(column: str) -> str:
    from recet.fieldnames import recet_name

    return recet_name(column)


def _summary(problems: list, records: int, notes: list[str], rule_set: str, key: str) -> dict:
    rejected = {(p.table, p.row, p.record_id) for p in problems if p.severity == "reject"}
    warned = {(p.table, p.row, p.record_id) for p in problems if p.severity != "reject"} - rejected
    by_rule = defaultdict(lambda: [0, "", ""])
    for p in problems:
        by_rule[p.rule_id][0] += 1
        by_rule[p.rule_id][1:] = [p.severity, p.message]
    top = sorted(by_rule.items(), key=lambda kv: (kv[1][1] != "reject", -kv[1][0]))
    return {
        "records": records,
        "rejected": len(rejected),
        "warnings": len(warned),
        "notes": notes,
        "rule_set": rule_set,
        "findings_key": key,
        "findings": len(problems),
        "by_rule": [{"rule_id": k, "count": c, "severity": s, "example": m} for k, (c, s, m) in top[:20]],
    }


@transaction.atomic
def upload(actor, deadline: Deadline, organisation, data: bytes, filename: str) -> SubmissionVersion:
    """Take a PA's upload for a filing: start the submission (or use its open draft), keep the job (the version's earlier file is
    deleted: superseded uploads are not kept, BR-04.4), read its budget lines, run the module's checks and the budget filing rule set,
    and record the result: validated when nothing is rejected."""
    if not hasattr(deadline, "budget_setup"):
        raise AccessDenied("This filing is not set up for budgets.")
    tables, job = jobfile.read_upload(data, filename)
    v = submissions.start(actor, deadline, organisation)
    lines, problems, records = derive(tables, organisation, deadline.cycle)
    found, notes, rule_set = _rule_findings(tables, deadline.cycle)
    problems += found
    store = storage.store("uploads")
    sha = hashlib.sha256(job).hexdigest()
    folder = f"{_prefix(v)}{sha[:12]}/"
    key, findings_key = folder + "job.zip", folder + "findings.json"
    store.put_bytes(key, job)
    store.put_bytes(findings_key, json.dumps([p.__dict__ for p in problems]).encode())
    old = {f["key"].rsplit("/", 1)[0] + "/" for f in v.files or [] if f.get("key")} - {folder}
    transaction.on_commit(lambda: [store.delete_prefix(o) for o in old])  # the superseded upload, once this one is recorded
    name = filename if filename.lower().endswith(".zip") else Path(filename).stem + ".zip"
    v = submissions.attach_files(actor, v, [{"name": name, "key": key, "sha256": sha, "bytes": len(job), "uploaded_as": filename}])
    _replace_lines(v, lines)
    summary = _summary(problems, records, notes, rule_set, findings_key)
    rs = rules.in_force(RULE_SET)
    return submissions.record_validation(actor, v, rs, summary)


def _replace_lines(v: SubmissionVersion, lines: list[dict]) -> None:
    BudgetLine.objects.filter(version=v).delete()
    BudgetLine.objects.bulk_create(BudgetLine(version=v, **line) for line in lines)
    BudgetResult.objects.filter(version=v).delete()


def job_file(v: SubmissionVersion) -> dict | None:
    return next((f for f in v.files or [] if f.get("key", "").endswith("job.zip")), None)


def ensure_lines(v: SubmissionVersion) -> None:
    """A reopened version (n+1) starts with its predecessor's files: read its lines from them."""
    f = job_file(v)
    if f is None or v.budget_lines.exists():
        return
    tables, _ = jobfile.read_upload(storage.store("uploads").get_bytes(f["key"]), "job.zip")
    lines, _, _ = derive(tables, v.submission.organisation, v.submission.deadline.cycle)
    _replace_lines(v, lines)


def findings(v: SubmissionVersion) -> list[dict]:
    key = (v.findings or {}).get("findings_key")
    if not key or not storage.store("uploads").exists(key):
        return []
    return json.loads(storage.store("uploads").get_bytes(key))


def findings_csv(v: SubmissionVersion) -> str:
    return rules.findings_csv([rules.Finding(**f) for f in findings(v)])


# --- pricing and confirming


def price(actor, v: SubmissionVersion) -> str:
    """Price a validated version on its filing's ACC (BR-02.5); a version with statewide programs is allocated with the filing's
    allocation version, or the most recent confirmed one for the program year, and the run records which (BR-04.8). Returns the run's
    id."""
    import tempfile

    setup = v.submission.deadline.budget_setup
    f = job_file(v)
    if f is None:
        raise AccessDenied("Upload the filing first.")
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "job.zip"
        path.write_bytes(storage.store("uploads").get_bytes(f["key"]))
        name = f"{v.submission.organisation.code} {label(v.submission.deadline)}, version {v.number}"
        statewide = Program.objects.filter(prg_id__in=v.budget_lines.values("prg_id"), statewide=True).exists()
        return submissions.price(
            actor,
            v,
            path,
            name=name,
            acc_version=setup.acc_version,
            allocation_version=setup.allocation_version,
            settings={"allocate_statewide": statewide},
        )


def priced(v: SubmissionVersion) -> bool:
    """A finished run on exactly the version's current job."""
    f = job_file(v)
    return bool(v.calculation and f and v.calculation.get("job_sha256") == f["sha256"])


def import_results(v: SubmissionVersion) -> bool:
    """Read the version's priced results into BudgetResult rows (once per run): the run's group sums, summed by credited PA, program and
    program year. False when the version has none (not priced, or the run kept no results)."""
    import pyarrow as pa
    import pyarrow.parquet as pq

    calc = v.calculation or {}
    key = (calc.get("results") or {}).get("group_sums")
    if not key:
        return False
    if BudgetResult.objects.filter(version=v, run=calc["run"]).exists():
        return True
    table = pq.read_table(pa.BufferReader(storage.store("results").get_bytes(key)))
    sums: dict[tuple, dict] = {}
    for row in table.to_pylist():
        year = row.get("_year")
        k = (str(row.get("_pa") or ""), str(row.get("_prg") or ""), int(year) if str(year or "").isdigit() else 0)
        s = sums.setdefault(k, {"name": "", "n": 0, "values": defaultdict(float)})
        s["name"] = s["name"] or str(row.get("_prgname") or "")
        s["n"] += int(row.get("_n") or 0)
        for c, val in row.items():
            if c not in RESULT_KEYS and isinstance(val, int | float) and val == val:  # NaN != NaN: a blank sum
                s["values"][c] += float(val)
    with transaction.atomic():
        BudgetResult.objects.filter(version=v).delete()
        BudgetResult.objects.bulk_create(
            BudgetResult(
                version=v,
                run=calc["run"],
                pa=k[0],
                prg_id=k[1],
                year=k[2],
                program_name=s["name"],
                measures=s["n"],
                values=dict(s["values"]),
            )
            for k, s in sums.items()
        )
    return True


def confirm(actor, v: SubmissionVersion) -> SubmissionVersion:
    """Confirm the priced version: the sign-off on exactly the results reviewed; publication follows with no further calculation."""
    if not priced(v):
        raise AccessDenied("Price this version first: confirming needs a finished run on exactly this upload.")
    return submissions.confirm(actor, v, v.calculation["input_hash"])


def on_publish(v: SubmissionVersion) -> None:
    """Publication hook: the results the version was confirmed on are read now, so the published version holds them."""
    import_results(v)


def on_unpublish(v: SubmissionVersion) -> None:
    """Un-confirmed before the due date: nothing is kept apart from the version, which goes back to validated."""


# --- results by program


METRICS = [
    ("budget", "Budget", "money", "Program costs plus measure rebates and incentives, nominal dollars (CET's Budget)"),
    ("TotalSystemBenefit", "Total system benefit", "money", "Total system benefit (TSB), present value"),
    ("TRC", "TRC", "ratio", "Total Resource Cost ratio: TRC benefits / TRC cost"),
    ("PAC", "PAC", "ratio", "Program Administrator Cost ratio: PAC benefits / PAC cost"),
    ("RIM", "RIM", "ratio", "Ratepayer Impact Measure ratio"),
    ("SCB", "SCT Base", "ratio", "Societal Cost Test, base social cost of carbon"),
    ("FirstYearNetKWh", "First-year net kWh", "number", "First-year net electric savings"),
    ("FirstYearNetKW", "First-year net kW", "number", "First-year net peak demand savings"),
    ("FirstYearNetThm", "First-year net therms", "number", "First-year net gas savings"),
    ("LifecycleNetKWh", "Lifecycle net kWh", "number", "Net electric savings over the measures' lives"),
    ("LifecycleNetThm", "Lifecycle net therms", "number", "Net gas savings over the measures' lives"),
]
RATIOS = {"TRC": ("TRCBen", "TRCCost"), "PAC": ("PACBen", "PACCost"), "RIM": ("RIMBen", "RIMCost"), "SCB": ("SCBBen", "SCBCost")}


@dataclass
class Row:
    prg_id: str
    name: str = ""
    pa: str = ""
    portfolio_cost: bool = False
    excluded_budget: bool = False
    excluded_ce: bool = False
    program_costs: Decimal = Decimal(0)
    rebates: Decimal = Decimal(0)
    on_bill_financing: Decimal = Decimal(0)
    by_field: dict = field(default_factory=lambda: defaultdict(Decimal))
    measures: int = 0
    sums: dict = field(default_factory=lambda: defaultdict(float))
    priced: bool = False

    @property
    def budget(self) -> Decimal:
        return self.program_costs + self.rebates

    def value(self, key: str):
        if key == "budget":
            return self.budget
        if key in RATIOS:
            ben, cost = (self.sums.get(c) for c in RATIOS[key])
            return ben / cost if ben is not None and cost else None
        return self.sums.get(key) if self.priced else None

    def add(self, other: "Row") -> None:
        self.program_costs += other.program_costs
        self.rebates += other.rebates
        self.on_bill_financing += other.on_bill_financing
        self.measures += other.measures
        for k, val in other.by_field.items():
            self.by_field[k] += val
        for k, val in other.sums.items():
            self.sums[k] += val
        self.priced = self.priced or other.priced


@dataclass
class Table:
    rows: list[Row]
    total: Row
    years: list[int]
    all_years: list[int]
    pas: list[str]
    reconciliation: dict
    priced: bool


def table(v: SubmissionVersion, *, years=None, pa: str = "", hide_budget_excluded=False, hide_ce_excluded=False) -> Table:
    """The version by program over the chosen program years (BR-04.1; none = every year): its filed budget from its own lines and its
    priced results summed over the years and credited PAs (or one credited PA). Programs flagged exclude-from-budget or exclude-from-C/E
    in any chosen year are marked and, when asked, left out of the rows and the total; their values are never changed (BR-04.3). The
    reconciliation always covers every line: the filed total, the part excluded from the budget, and the total without it."""
    ensure_lines(v)
    import_results(v)
    lines = list(v.budget_lines.all())
    results = list(v.budget_results.filter(run=(v.calculation or {}).get("run", "")))
    all_years = sorted({ln.year for ln in lines} | {r.year for r in results if r.year})
    chosen = sorted(int(y) for y in years or [] if str(y).isdigit() and int(y) in all_years) or all_years
    pas = sorted({r.pa for r in results if r.pa})
    rows: dict[str, Row] = {}
    for ln in lines:
        if ln.year not in chosen:
            continue
        r = rows.setdefault(ln.prg_id, Row(ln.prg_id))
        r.pa = r.pa or ln.pa
        costs = {k: Decimal(val) for k, val in (ln.program_costs or {}).items()}
        for k, val in costs.items():
            r.by_field[k] += val
        r.program_costs += sum(costs.values(), Decimal(0))
        r.rebates += ln.rebates_and_incentives
        r.on_bill_financing += ln.on_bill_financing
        r.measures += ln.measures
    for res in results:
        if (res.year and res.year not in chosen) or (pa and res.pa != pa):
            continue
        r = rows.setdefault(res.prg_id, Row(res.prg_id))
        r.name = r.name or res.program_name
        r.priced = True
        for k, val in res.values.items():
            r.sums[k] += val
    programs = {p.prg_id: p for p in Program.objects.filter(prg_id__in=rows).prefetch_related("versions")}
    for r in rows.values():
        p = programs.get(r.prg_id)
        if p is None:
            continue
        r.portfolio_cost = p.portfolio_cost
        versions = sorted(p.versions.all(), key=lambda x: x.valid_from, reverse=True)
        for y in chosen:
            pv = registry._for_year(versions, y)
            if pv is None:
                continue
            r.name = pv.name if not r.name or y == chosen[0] else r.name
            r.excluded_budget |= pv.exclude_from_budget
            r.excluded_ce |= pv.exclude_from_ce
    everything = sorted(rows.values(), key=lambda r: r.prg_id)
    shown = [r for r in everything if not (hide_budget_excluded and r.excluded_budget) and not (hide_ce_excluded and r.excluded_ce)]
    total = Row("Total")
    for r in shown:
        total.add(r)
    filed = sum((r.budget for r in everything), Decimal(0))
    excluded = sum((r.budget for r in everything if r.excluded_budget), Decimal(0))
    reconciliation = {
        "filed": filed,
        "excluded_budget": excluded,
        "without_excluded": filed - excluded,
        "on_bill_financing": sum((r.on_bill_financing for r in everything), Decimal(0)),
        "portfolio_costs": sum((r.budget for r in everything if r.portfolio_cost), Decimal(0)),
    }
    return Table(shown, total, chosen if chosen != all_years else [], all_years, pas, reconciliation, bool(results))


def compare(a: SubmissionVersion, b: SubmissionVersion, *, years=None, changed_only=False) -> list[dict]:
    """Two versions side by side by program (BR-04.2): each metric in a, in b and the difference; programs in one version only are
    blank for the other. Neither version is changed."""
    ta, tb = table(a, years=years), table(b, years=years)
    ra, rb = {r.prg_id: r for r in ta.rows}, {r.prg_id: r for r in tb.rows}
    out = []
    for prg in sorted(set(ra) | set(rb)):
        x, y = ra.get(prg), rb.get(prg)
        cells, changed = [], False
        for key, *_ in METRICS:
            va, vb = (x.value(key) if x else None), (y.value(key) if y else None)
            diff = float(vb) - float(va) if va is not None and vb is not None else None
            changed |= (va is None) != (vb is None) or (diff is not None and abs(diff) > 1e-9 * max(1.0, abs(float(va))))
            cells.append((va, vb, diff))
        if changed or not changed_only:
            out.append(
                {"prg_id": prg, "name": (x or y).name, "cells": cells, "changed": changed, "only": "a" if not y else "b" if not x else ""}
            )
    totals = []
    for key, *_ in METRICS:
        va, vb = ta.total.value(key), tb.total.value(key)
        totals.append((va, vb, float(vb) - float(va) if va is not None and vb is not None else None))
    out.append({"prg_id": "Total", "name": "", "cells": totals, "changed": True, "only": "", "total": True})
    return out


# --- published inputs and the appendix


def published_files(v: SubmissionVersion) -> list[tuple[str, str, str]]:
    """(label, area, key) of what a confirmed version publishes so it can be reproduced (BR-04.4): its job, the run's job spec (the
    settings, release, ACC, packs and allocation table it was priced with) and the run's receipt."""
    calc, out = v.calculation or {}, []
    f = job_file(v)
    if f:
        out.append(("Inputs (Measure and ProgramCost)", "uploads", f["key"]))
    if calc.get("spec_key"):
        out.append(("Run settings (job spec)", "results", calc["spec_key"]))
    if calc.get("receipt_key"):
        out.append(("Run receipt (engine, reference data, outputs)", "results", calc["receipt_key"]))
    return out


def appendix_xlsx(v: SubmissionVersion) -> bytes:
    """The filing appendix tables generated from the version's data (BR-04.5): budget by program and year with the cost categories,
    cost-effectiveness by program, savings by program and year, and what the version was priced with."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    about = wb.active
    about.title = "About"
    calc = v.calculation or {}
    deadline = v.submission.deadline
    info = [
        ("Filing", label(deadline)),
        ("Program administrator", v.submission.organisation.code),
        ("Version", v.number),
        ("Status", v.get_status_display()),
        ("Confirmed", v.confirmed_at.isoformat() if v.confirmed_at else ""),
        ("ACC", f"{calc.get('acc')} ACC" if calc.get("acc") else ""),
        ("Statewide allocation version", (calc.get("allocation_version") or {}).get("label", "")),
        ("Engine release", calc.get("release", "")),
        ("Calculation run", calc.get("run", "")),
        ("Rule set", (v.findings or {}).get("rule_set", "")),
    ]
    for k, val in info:
        about.append([k, val])
    about["A1"].font = Font(bold=True)
    lines = sorted(v.budget_lines.all(), key=lambda ln: (ln.prg_id, ln.year))
    ws = wb.create_sheet("Budget")
    ws.append(
        ["PrgID", "Year", "PA"] + [f for f, _ in PROGRAM_COST_FIELDS] + ["ProgramCosts", "RebatesandIncents", "Budget", "OnBillFinancing"]
    )
    for ln in lines:
        costs = [float(Decimal(ln.program_costs.get(f, "0"))) for f, _ in PROGRAM_COST_FIELDS]
        ws.append(
            [
                ln.prg_id,
                ln.year,
                ln.pa,
                *costs,
                sum(costs),
                float(ln.rebates_and_incentives),
                sum(costs) + float(ln.rebates_and_incentives),
                float(ln.on_bill_financing),
            ]
        )
    t = table(v)
    ws = wb.create_sheet("Cost-effectiveness")
    ws.append(["PrgID", "Program", "Excluded from budget", "Excluded from C/E"] + [m[1] for m in METRICS])
    for r in t.rows + [t.total]:
        ws.append(
            [r.prg_id, r.name, "yes" if r.excluded_budget else "", "yes" if r.excluded_ce else ""]
            + [None if r.value(k) is None else float(r.value(k)) for k, *_ in METRICS]
        )
    ws = wb.create_sheet("Savings by year")
    keys = ["FirstYearNetKWh", "FirstYearNetKW", "FirstYearNetThm", "LifecycleNetKWh", "LifecycleNetThm", "TotalSystemBenefit"]
    ws.append(["PrgID", "Year", "Credited PA"] + keys)
    for res in v.budget_results.filter(run=calc.get("run", "")).order_by("prg_id", "year", "pa"):
        ws.append([res.prg_id, res.year or None, res.pa] + [res.values.get(k) for k in keys])
    for sheet in wb.worksheets[1:]:
        for c in sheet[1]:
            c.font = Font(bold=True)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
