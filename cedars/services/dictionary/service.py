"""The field catalog and what is generated from it. Every function reads the current (newest) catalog version unless given one."""

import csv
import io
from dataclasses import dataclass, field

from django.db import transaction

from cedars.services import audit
from cedars.services.accounts import service as access
from cedars.services.accounts.service import AccessDenied

from .models import Alias, CatalogVersion, Field, Usage

ATTRIBUTES = [
    "unit",
    "source",
    "source_version",
    "cost_tests",
    "time_basis",
    "aggregation",
    "rounding",
    "years_in_use",
    "audience",
    "required_by",
]
JSON_TYPES = {"text": "string", "number": "number", "boolean": "boolean", "year_quarter": "string"}


def current() -> CatalogVersion | None:
    return CatalogVersion.objects.first()


@transaction.atomic
def load_catalog(actor, payload: dict, *, label: str | None = None, notes: str = "") -> CatalogVersion:
    """A new catalog version from an exported catalog (scripts/export_recet_dictionary.py writes one from ReCET)."""
    if actor is not None and not access.is_system_admin(actor):
        raise AccessDenied("Only a CPUC system administrator publishes a dictionary version.")
    label = label or payload.get("source") or "catalog"
    version = CatalogVersion.objects.create(label=label, notes=notes, loaded_by=actor)
    for f in payload["fields"]:
        obj = Field.objects.create(
            version=version,
            name=f["name"],
            kind=f["kind"],
            group=f.get("group", ""),
            definition=f.get("definition", ""),
            doc_link=f.get("doc_link", ""),
            data_type=f.get("data_type", "text"),
            max_length=f.get("max_length"),
            unit=f.get("unit", ""),
            source=f.get("source", ""),
            source_version=f.get("version", ""),
            cost_tests=f.get("cost_tests", ""),
            time_basis=f.get("time_basis", ""),
            aggregation=f.get("aggregation", ""),
            rounding=f.get("rounding", ""),
            years_in_use=f.get("years_in_use", ""),
            audience=f.get("audience", Field.Audience.PUBLIC),
            required_by=f.get("required_by", ""),
        )
        Alias.objects.bulk_create(Alias(field=obj, name=a["name"], system=a["system"]) for a in f.get("aliases", []))
        Usage.objects.bulk_create(
            Usage(field=obj, module=u["module"], table=u["table"], position=u["position"], required=u.get("required", False))
            for u in f.get("usage", [])
        )
    audit.record("dictionary.published", obj=version, after={"label": label, "fields": len(payload["fields"])}, actor=actor)
    return version


def resolve(name: str, version: CatalogVersion | None = None) -> Field | None:
    """The field a column name means: its canonical name or any alias, ignoring case and surrounding spaces."""
    version = version or current()
    if version is None:
        return None
    n = str(name).strip()
    return Field.objects.filter(version=version, name__iexact=n).first() or (
        Field.objects.filter(version=version, aliases__name__iexact=n).first()
    )


def table_fields(module: str, table: str, version: CatalogVersion | None = None) -> list[tuple[Field, Usage]]:
    version = version or current()
    usages = Usage.objects.filter(field__version=version, module=module, table=table).select_related("field").order_by("position")
    return [(u.field, u) for u in usages]


@dataclass
class HeaderMap:
    """How an uploaded file's columns line up with a module's table: `columns` maps each header to its canonical name; `renamed` are
    headers given under another name (CET's), `unknown` are not in the catalog (ignored, not rejected), `not_used` are catalog fields
    this table does not use (ignored too), `missing` are required fields the file lacks."""

    columns: dict[str, str] = field(default_factory=dict)
    renamed: dict[str, str] = field(default_factory=dict)
    unknown: list[str] = field(default_factory=list)
    not_used: list[str] = field(default_factory=list)
    missing: list[str] = field(default_factory=list)


def map_headers(headers: list[str], module: str, table: str, version: CatalogVersion | None = None) -> HeaderMap:
    version = version or current()
    wanted = {f.name: u for f, u in table_fields(module, table, version)}
    lookup = {}
    for f in Field.objects.filter(version=version).prefetch_related("aliases"):
        lookup[f.name.lower()] = f.name
        for a in f.aliases.all():
            lookup.setdefault(a.name.lower(), f.name)
    out = HeaderMap()
    for h in headers:
        canonical = lookup.get(str(h).strip().lower())
        if canonical is None:
            out.unknown.append(h)
        elif canonical not in wanted:
            out.not_used.append(h)
        else:
            out.columns[h] = canonical
            if str(h).strip() != canonical:
                out.renamed[h] = canonical
    present = set(out.columns.values())
    out.missing = [name for name, u in wanted.items() if u.required and name not in present]
    return out


def template_csv(module: str, table: str, version: CatalogVersion | None = None) -> str:
    """An empty upload file: the table's header row in column order (pipe-delimited, as CET's files are)."""
    return "|".join(f.name for f, _ in table_fields(module, table, version)) + "\n"


def template_xlsx(module: str, table: str, version: CatalogVersion | None = None) -> bytes:
    """A workbook with the table's columns on its first sheet and a Fields sheet defining each one."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    version = version or current()
    rows = table_fields(module, table, version)
    wb = Workbook()
    ws = wb.active
    ws.title = table[:31]
    ws.append([f.name for f, _ in rows])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    doc = wb.create_sheet("Fields")
    doc.append(["Field", "Required", "Type", "Unit", "Definition", "Also accepted as", f"Catalog: {version.label}"])
    for f, u in rows:
        doc.append([f.name, "yes" if u.required else "", f.data_type, f.unit, f.definition, ", ".join(a.name for a in f.aliases.all())])
    for cell in doc[1]:
        cell.font = Font(bold=True)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def json_schema(module: str, table: str, version: CatalogVersion | None = None) -> dict:
    """A JSON Schema for one row of the table, for the API (a row is an object keyed by canonical names)."""
    version = version or current()
    props, required = {}, []
    for f, u in table_fields(module, table, version):
        p = {"type": [JSON_TYPES.get(f.data_type, "string"), "null"], "description": f.definition}
        if f.max_length:
            p["maxLength"] = f.max_length
        if f.data_type == "year_quarter":
            p["pattern"] = r"^\d{4}Q[1-4]$"
        if f.unit:
            p["x-unit"] = f.unit
        props[f.name] = p
        if u.required:
            required.append(f.name)
    return {
        "$schema": "https://json-schema.org/draft/2020-12/schema",
        "title": f"{module} {table} row",
        "description": f"Generated from the CEDARS field catalog {version.label}",
        "type": "object",
        "properties": props,
        "required": required,
    }


def catalog_csv(version: CatalogVersion | None = None) -> str:
    version = version or current()
    out = io.StringIO()
    w = csv.writer(out, lineterminator="\n")
    w.writerow(["name", "kind", "group", "definition", "data_type", "max_length", *ATTRIBUTES, "also_known_as", "documentation"])
    for f in Field.objects.filter(version=version).prefetch_related("aliases").order_by("kind", "group", "name"):
        w.writerow(
            [
                f.name,
                f.kind,
                f.group,
                f.definition,
                f.data_type,
                f.max_length or "",
                *[getattr(f, a) for a in ATTRIBUTES],
                "; ".join(f"{a.name} ({a.system})" for a in f.aliases.all()),
                f.doc_link,
            ]
        )
    return out.getvalue()


def crosswalk_csv(version: CatalogVersion | None = None) -> str:
    """Every other name and the canonical name it maps to."""
    version = version or current()
    out = io.StringIO()
    w = csv.writer(out, lineterminator="\n")
    w.writerow(["other_name", "used_in", "canonical_name"])
    for a in Alias.objects.filter(field__version=version).select_related("field").order_by("name"):
        w.writerow([a.name, a.system, a.field.name])
    return out.getvalue()
